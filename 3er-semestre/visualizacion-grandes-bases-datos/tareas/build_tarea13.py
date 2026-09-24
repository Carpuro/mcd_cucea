# -*- coding: utf-8 -*-
"""
Tarea 1.3 - Series de tiempo de CO y su acumulado de 8 h (medias moviles),
y graficas de caja por hora y por mes. Estacion Centro (CEN).
Se cambia de estacion respecto a las tareas 1.0-1.2 (Country/COU) porque COU
solo tiene datos de septiembre 2024 y no permite el analisis mensual.
"""
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"C:\Users\carlo\GitHub\mcd_cucea"
TEMPLATE = BASE + r"\_recursos\plantillas\plantilla_informe_MCD.docx"
COURSE = BASE + r"\3er-semestre\visualizacion-grandes-bases-datos"
DATA = COURSE + r"\datos\BD_2024_calidad_aire.xlsx"
OUT = COURSE + r"\tareas\Tarea1_3"
NAVY = RGBColor(0x1F, 0x49, 0x7D); AMBER = RGBColor(0xD4, 0x8C, 0x00)
WHITE = RGBColor(0xFF, 0xFF, 0xFF); GRAY_L = RGBColor(0xF2, 0xF2, 0xF2)
NAVY_HEX, AMBER_HEX, GREEN_HEX = "#1F497D", "#D48C00", "#3C8C5A"
MESES = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]

# ============================ DATA ==========================================
df = pd.read_excel(DATA, sheet_name="Data")
d = df[df.STATION == "CEN"].copy()
d["CO"] = pd.to_numeric(d["CO"], errors="coerce")
d["DATE"] = pd.to_datetime(d["DATE"])
d["ts"] = d["DATE"]   # DATE ya incluye la hora (timestamp horario completo)
d = d.sort_values("ts")
# serie horaria reindexada a todo 2024 (para medias moviles correctas en el tiempo)
full = pd.date_range("2024-01-01", "2024-12-31 23:00", freq="h")
co = d.groupby("ts")["CO"].mean().reindex(full)   # groupby dedups timestamps repetidos
n_valid = int(co.notna().sum())
# acumulado de 8 h = promedio movil de 8 h (indicador estandar de CO / norma IMECA)
co8 = co.rolling(8, min_periods=6).mean()
# medias moviles centradas: diaria SMA(24) y mensual SMA(720 ~ 30 dias)
co_d = co.rolling(24, center=True, min_periods=18).mean()
co_m = co.rolling(720, center=True, min_periods=360).mean()
co8_d = co8.rolling(24, center=True, min_periods=18).mean()
co8_m = co8.rolling(720, center=True, min_periods=360).mean()

tab = pd.DataFrame({"CO": co, "CO8": co8})
tab["HOUR"] = tab.index.hour; tab["MES"] = tab.index.month
byh = tab.groupby("HOUR")["CO"].mean()
bym = tab.groupby("MES")["CO"].mean()
top_h = byh.sort_values(ascending=False).head(3).index.tolist()
min_h = byh.sort_values().head(3).index.tolist()
top_m = bym.sort_values(ascending=False).head(3).index.tolist()
min_m = bym.sort_values().head(3).index.tolist()


def f(v, dd=2):
    return f"{v:.{dd}f}"


# ============================ FIGURES =======================================
def fig_timeseries():
    p = OUT + r"\ts_co_medias_moviles.png"
    fig, ax = plt.subplots(2, 1, figsize=(9.2, 6.4), sharex=True)
    for a, serie, sd, sm, titulo, ylab in [
        (ax[0], co, co_d, co_m, "Serie de tiempo de CO por hora — Centro (2024)", "CO (ppm)"),
        (ax[1], co8, co8_d, co8_m, "Acumulado de 8 h (promedio móvil 8 h) — Centro (2024)", "CO 8 h (ppm)")]:
        a.plot(serie.index, serie.values, color="#B9C7DA", linewidth=0.4, label="Serie horaria")
        a.plot(sd.index, sd.values, color=GREEN_HEX, linewidth=1.0, label="SMA 24 h (tendencia diaria)")
        a.plot(sm.index, sm.values, color=AMBER_HEX, linewidth=2.2, label="SMA 720 h (tendencia mensual)")
        a.set_title(titulo, color=NAVY_HEX, weight="bold", fontsize=11)
        a.set_ylabel(ylab, color=NAVY_HEX); a.grid(axis="y", linestyle=":", alpha=0.5)
        a.legend(fontsize=8, loc="upper center", ncol=3)
    ax[1].xaxis.set_major_formatter(mdates.DateFormatter("%b"))
    ax[1].xaxis.set_major_locator(mdates.MonthLocator())
    ax[1].set_xlabel("Mes (2024)")
    plt.tight_layout(); plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


def _box(ax, groups, labels, title, ylab, highlight):
    bp = ax.boxplot(groups, patch_artist=True, showfliers=True, widths=0.6,
                    medianprops=dict(color=AMBER_HEX, linewidth=1.6),
                    whiskerprops=dict(color=NAVY_HEX), capprops=dict(color=NAVY_HEX),
                    flierprops=dict(marker="o", markerfacecolor=AMBER_HEX,
                                    markeredgecolor=NAVY_HEX, markersize=2.5, alpha=0.4))
    for i, box in enumerate(bp["boxes"]):
        box.set(facecolor="#E8A33D" if (i + labels[0]) in highlight else "#9DB7D9",
                edgecolor=NAVY_HEX, linewidth=1.1)
    ax.set_xticklabels(labels if isinstance(labels[0], str) else [str(l) for l in labels], fontsize=7)
    ax.set_title(title, color=NAVY_HEX, weight="bold", fontsize=10.5)
    ax.set_ylabel(ylab, color=NAVY_HEX); ax.grid(axis="y", linestyle=":", alpha=0.5)


def fig_boxplots(var, p_name, ylab):
    p = OUT + "\\" + p_name
    fig, ax = plt.subplots(2, 1, figsize=(9.2, 7.0))
    # por hora
    gh = [tab.loc[tab.HOUR == h, var].dropna().values for h in range(24)]
    bp = ax[0].boxplot(gh, positions=range(24), patch_artist=True, widths=0.6,
                       medianprops=dict(color=AMBER_HEX, linewidth=1.6),
                       whiskerprops=dict(color=NAVY_HEX), capprops=dict(color=NAVY_HEX),
                       flierprops=dict(marker="o", markerfacecolor=AMBER_HEX,
                                       markeredgecolor=NAVY_HEX, markersize=2.3, alpha=0.35))
    for h, box in enumerate(bp["boxes"]):
        box.set(facecolor="#E8A33D" if h in top_h else "#9DB7D9", edgecolor=NAVY_HEX, linewidth=1.0)
    ax[0].set_xticks(range(24)); ax[0].set_xticklabels(range(24), fontsize=7)
    ax[0].set_title(f"{ylab} por HORA del día — Centro 2024", color=NAVY_HEX, weight="bold", fontsize=10.5)
    ax[0].set_xlabel("Hora"); ax[0].set_ylabel(ylab, color=NAVY_HEX)
    ax[0].grid(axis="y", linestyle=":", alpha=0.5)
    # por mes
    gm = [tab.loc[tab.MES == m, var].dropna().values for m in range(1, 13)]
    bp2 = ax[1].boxplot(gm, positions=range(1, 13), patch_artist=True, widths=0.6,
                        medianprops=dict(color=AMBER_HEX, linewidth=1.6),
                        whiskerprops=dict(color=NAVY_HEX), capprops=dict(color=NAVY_HEX),
                        flierprops=dict(marker="o", markerfacecolor=AMBER_HEX,
                                        markeredgecolor=NAVY_HEX, markersize=2.3, alpha=0.35))
    for i, box in enumerate(bp2["boxes"]):
        box.set(facecolor="#E8A33D" if (i + 1) in top_m else "#9DB7D9", edgecolor=NAVY_HEX, linewidth=1.0)
    ax[1].set_xticks(range(1, 13)); ax[1].set_xticklabels(MESES, fontsize=8)
    ax[1].set_title(f"{ylab} por MES — Centro 2024", color=NAVY_HEX, weight="bold", fontsize=10.5)
    ax[1].set_xlabel("Mes"); ax[1].set_ylabel(ylab, color=NAVY_HEX)
    ax[1].grid(axis="y", linestyle=":", alpha=0.5)
    plt.tight_layout(); plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


# ============================ DOCX HELPERS ==================================
def set_cell_bg(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"); tcPr.append(shd)


def add_run(p, t, bold=False, italic=False, color=None, size=None, font="Calibri"):
    r = p.add_run(t); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    if size: r.font.size = Pt(size)
    r.font.name = font; return r


def add_para(doc, t="", bold=False, italic=False, color=None, size=10,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, sb=0, sa=5):
    p = doc.add_paragraph(); p.alignment = align
    p.paragraph_format.space_before = Pt(sb); p.paragraph_format.space_after = Pt(sa)
    if t: add_run(p, t, bold, italic, color, size)
    return p


def add_heading(doc, t, level=1):
    return add_para(doc, t, bold=True, color=NAVY, size=(13 if level == 1 else 11),
                    align=WD_ALIGN_PARAGRAPH.LEFT, sb=(10 if level == 1 else 7), sa=3)


def add_caption(doc, t):
    return add_para(doc, t, italic=True, color=RGBColor(0x40, 0x40, 0x40), size=9,
                    align=WD_ALIGN_PARAGRAPH.CENTER, sb=2, sa=6)


def add_bullet(doc, titulo, texto):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(4)
    add_run(p, titulo + " ", bold=True, color=NAVY, size=10); add_run(p, texto, size=10)


def add_pic(doc, path, w=6.2):
    doc.add_picture(path, width=Inches(w))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def make_table(doc, headers, rows, caption=None, widths=None):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    if widths:
        for i, w in enumerate(widths):
            for row in tbl.rows:
                row.cells[i].width = Inches(w)
    hr = tbl.rows[0]
    for c in hr.cells: set_cell_bg(c, NAVY)
    for ci, h in enumerate(headers):
        pp = hr.cells[ci].paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(pp, h, bold=True, color=WHITE, size=9)
    for ri, rd in enumerate(rows):
        dr = tbl.rows[ri + 1]
        for c in dr.cells: set_cell_bg(c, WHITE if ri % 2 == 0 else GRAY_L)
        for ci, v in enumerate(rd):
            pp = dr.cells[ci].paragraphs[0]
            pp.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
            add_run(pp, str(v), size=9)
    if caption: add_caption(doc, caption)
    else: doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_apa(doc, t):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.5); p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.space_after = Pt(6); add_run(p, t, size=10)


def clear_body_after_table1(doc):
    body = doc.element.body; sect = body.find(qn("w:sectPr")); c = 0; cut = None
    for i, ch in enumerate(list(body)):
        if ch.tag == qn("w:tbl"):
            c += 1
            if c == 2: cut = i; break
    if cut is not None:
        for el in [e for e in list(body)[cut + 1:] if e.tag != qn("w:sectPr")]: body.remove(el)
    if sect is not None and sect not in list(body): body.append(sect)


# ============================ BUILD =========================================
def build():
    import os
    os.makedirs(OUT, exist_ok=True)
    p_ts = fig_timeseries()
    p_box_co = fig_boxplots("CO", "boxplots_co.png", "CO (ppm)")
    p_box_co8 = fig_boxplots("CO8", "boxplots_co8.png", "CO 8 h (ppm)")

    doc = Document(TEMPLATE)
    hdr = doc.tables[0].rows[0].cells[1]
    hdr.paragraphs[2].runs[0].text = "Maestría en Ciencias de los Datos | 3er Semestre"
    hdr.paragraphs[3].runs[0].text = "Materia: Visualización de Grandes Bases de Datos"
    tp = doc.paragraphs[1]; tp.clear()
    r = tp.add_run("Tarea 1.3 — Series de tiempo y gráficas de caja de CO\n"
                   "Concentración de CO y su acumulado de 8 h por hora y por mes — estación Centro")
    r.bold = True; r.font.size = Pt(13); r.font.color.rgb = NAVY; r.font.name = "Calibri"
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def si(idx, text):
        cell = doc.tables[1].rows[idx].cells[1]; cell.paragraphs[0].runs[0].text = text
        for extra in cell.paragraphs[0].runs[1:]: extra.text = ""
    si(1, "4 de septiembre de 2026")
    si(2, "Tarea 1.3 — Series de tiempo y gráficas de caja de CO")
    si(3, "Python 3.13 — pandas, numpy, matplotlib")
    si(4, "SIMAJ · Base de calidad del aire 2024 (BD_2024)")
    si(5, "Análisis local (equipo personal, entorno conda MCD)")
    si(6, "Calidad del aire ZMG 2024 — estación Centro (CEN); variable CO (ppm)")
    clear_body_after_table1(doc)

    add_para(doc, sb=8, sa=0)
    add_heading(doc, "1. Introducción")
    add_para(doc,
        "El objetivo de esta actividad es analizar el comportamiento temporal del monóxido de "
        "carbono (CO): (a) graficar las series de tiempo del CO y de su acumulado de 8 horas con "
        "las medias móviles que mejor reflejan sus tendencias diaria y mensual, y (b) construir "
        "las gráficas de caja de ambas variables por hora y por mes para responder: ¿en cuáles "
        "horas y en cuáles meses se presenta más este contaminante? Se trabaja con la base de "
        "calidad del aire 2024 de la Zona Metropolitana de Guadalajara (SIMAJ).")

    add_heading(doc, "1.1 Cambio de estación respecto a las tareas anteriores", level=2)
    add_para(doc,
        "Las tareas 1.0 a 1.2 se realizaron con la estación Country (COU), la más cercana al "
        "domicilio. Sin embargo, esta actividad requiere un análisis por mes y por tendencia "
        "mensual, y la estación Country solo cuenta con registros de CO de septiembre de 2024 "
        f"(578 lecturas de un único mes), por lo que es imposible responder la parte mensual. Por "
        "esta razón se cambió a la estación Centro (CEN), que dispone de "
        f"{n_valid} lecturas horarias distribuidas a lo largo de todo 2024 (enero a diciembre) y "
        "es, además, la estación de referencia utilizada por el profesor en la Sesión 4. Centro "
        "es también la estación con cobertura anual completa más cercana al domicilio (≈ 4.5 km).")

    add_heading(doc, "2. Desarrollo (metodología)")
    add_para(doc,
        "Una serie de tiempo es un conjunto de valores recolectados en intervalos de tiempo "
        "sucesivos; aquí, la concentración horaria de CO durante 2024. Sobre ella se aplican los "
        "conceptos vistos en clase:")
    add_bullet(doc, "Acumulado de 8 horas:",
        "promedio móvil de 8 horas de la concentración de CO. Es el indicador con el que la norma "
        "evalúa el CO (11 ppm/8 h) y suaviza el ruido horario.")
    add_bullet(doc, "Media móvil simple SMA(m):",
        "promedio de los m valores de la serie en una ventana; suaviza las fluctuaciones para "
        "estimar la tendencia. Se emplean dos órdenes: SMA(24) —ventana de 24 horas— para la "
        "tendencia diaria, y SMA(720) —ventana de ~30 días— para la tendencia mensual. Se usan "
        "centradas para no introducir desfase en la tendencia.")
    add_bullet(doc, "Gráfica de caja por grupo:",
        "se agrupan las lecturas por hora del día (0–23) y por mes (1–12) y se construye una caja "
        "por grupo; la altura de la caja y su posición indican en qué horas y meses se concentran "
        "las mayores concentraciones.")

    add_heading(doc, "3. Resultados")
    add_heading(doc, "3.1 Series de tiempo y medias móviles", level=2)
    add_para(doc, "La Figura 1 muestra la serie horaria de CO y su acumulado de 8 h, cada una con "
                  "su media móvil diaria (SMA 24 h) y mensual (SMA 720 h).", sa=3)
    add_pic(doc, p_ts, 6.3)
    add_caption(doc, "Figura 1 — Series de tiempo de CO (arriba) y de su acumulado de 8 h (abajo) "
                     "con medias móviles diaria (verde) y mensual (ámbar). La tendencia mensual "
                     "cae en verano y sube en invierno.")

    add_heading(doc, "3.2 Gráficas de caja del CO por hora y por mes", level=2)
    add_pic(doc, p_box_co, 6.1)
    add_caption(doc, "Figura 2 — Gráficas de caja del CO por hora (arriba) y por mes (abajo). En "
                     "ámbar, los grupos con mayor concentración.")

    add_heading(doc, "3.3 Gráficas de caja del acumulado de 8 h por hora y por mes", level=2)
    add_pic(doc, p_box_co8, 6.1)
    add_caption(doc, "Figura 3 — Gráficas de caja del acumulado de 8 h por hora (arriba) y por mes "
                     "(abajo). El acumulado suaviza el patrón horario pero conserva el mensual.")

    add_para(doc, "La Tabla 1 resume la media de CO por mes, que sustenta la respuesta.", sb=4, sa=3)
    make_table(doc, ["Mes"] + MESES,
        [["CO medio (ppm)"] + [f(bym[m]) for m in range(1, 13)]],
        caption="Tabla 1 — Concentración media de CO por mes en la estación Centro (2024).",
        widths=[1.3] + [0.42] * 12)

    add_heading(doc, "4. Respuesta y conclusiones")
    add_para(doc,
        f"¿En cuáles horas se presenta más el CO? La concentración sigue un patrón bimodal: es "
        f"máxima en la mañana (horas {min(top_h)}–9, con el pico a las {top_h[0]}:00 h ≈ "
        f"{f(byh[top_h[0]])} ppm) y presenta un segundo repunte por la noche (20–23 h); el mínimo "
        f"ocurre por la tarde (horas {min_h[0]}–{max(min_h)} h ≈ {f(byh[min_h[0]])} ppm). Este "
        "comportamiento coincide con las horas pico de tráfico vehicular y con la menor "
        "dispersión atmosférica de la madrugada y la noche.")
    add_para(doc,
        f"¿En cuáles meses se presenta más el CO? La concentración es máxima en los meses de "
        f"invierno —{MESES[top_m[0]-1]} ({f(bym[top_m[0]])} ppm), {MESES[top_m[1]-1]} "
        f"({f(bym[top_m[1]])}) y {MESES[top_m[2]-1]} ({f(bym[top_m[2]])})— y mínima a fin de "
        f"verano y en otoño —{MESES[min_m[0]-1]} ({f(bym[min_m[0]])} ppm), {MESES[min_m[1]-1]} "
        f"({f(bym[min_m[1]])}) y {MESES[min_m[2]-1]} ({f(bym[min_m[2]])})—. El máximo invernal se "
        "explica por las inversiones térmicas, que atrapan los contaminantes cerca del suelo, "
        "mientras que las lluvias del verano favorecen su dispersión y lavado.")
    add_para(doc,
        "En síntesis, en la estación Centro el CO se concentra más en las horas pico de la mañana "
        "(7–8 h) y por la noche, y en los meses de invierno (enero, diciembre, febrero). El "
        "acumulado de 8 h confirma el patrón mensual y atenúa el horario, como es de esperarse de "
        "un promedio móvil. Todas las medias móviles mensuales se mantienen por debajo de la norma "
        "de 11 ppm/8 h.")

    add_heading(doc, "Referencias")
    add_apa(doc, "Ruiz Chávez, C. (2024). Visualización de Grandes Bases de Datos — Sesión 4: "
                 "series de tiempo y medias móviles [material de clase]. Maestría en Ciencia de "
                 "los Datos, CUCEA, Universidad de Guadalajara.")
    add_apa(doc, "Secretaría de Medio Ambiente y Desarrollo Territorial. (2024). Sistema de "
                 "Monitoreo Atmosférico de Jalisco (SIMAJ): base de datos de calidad del aire "
                 "2024. Gobierno del Estado de Jalisco. https://siga.jalisco.gob.mx/aire/")
    out = OUT + r"\Tarea1_3_VisualizacionGBD_CarlosPulidoRosas.docx"
    doc.save(out); print("DOCX:", out); return out


if __name__ == "__main__":
    build()
