# -*- coding: utf-8 -*-
"""
Tarea 1.4 - Pollutant with the highest linear correlation with NO2 (station Centro).
Aligned with Session 6: covariance and Pearson r, scatter plots with regression lines,
linear vs non-linear fits (log, quadratic) and continuous vs discrete (logistic) models.
Country (COU) does not record NO2 in 2024, so Centro (CEN) is used.
"""
import os
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import statsmodels.api as sm
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"C:\Users\carlo\GitHub\mcd_cucea"
TEMPLATE = BASE + r"\_recursos\plantillas\plantilla_informe_MCD.docx"
COURSE = BASE + r"\3er-semestre\visualizacion-grandes-bases-datos"
DATA = COURSE + r"\datos\BD_2024_calidad_aire.xlsx"
OUT = COURSE + r"\tareas\Tarea1_4"
NAVY = RGBColor(0x1F, 0x49, 0x7D); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY_L = RGBColor(0xF2, 0xF2, 0xF2); HL = RGBColor(0xFC, 0xE4, 0xB8)
NAVY_HEX, AMBER_HEX = "#1F497D", "#D48C00"

CAND = ["O3", "NO", "NOX", "CO", "PM10", "PM2.5"]
NAMES = {"O3": "O₃ (ozono)", "NO": "NO (monóxido de nitrógeno)", "NOX": "NOx (óxidos de nitrógeno)",
         "CO": "CO (monóxido de carbono)", "PM10": "PM10 (partículas < 10 µm)",
         "PM2.5": "PM2.5 (partículas < 2.5 µm)"}
SHORT = {"O3": "O₃", "NO": "NO", "NOX": "NOx", "CO": "CO", "PM10": "PM10", "PM2.5": "PM2.5"}
UNIT = {"O3": "ppm", "NO": "ppm", "NOX": "ppm", "CO": "ppm", "PM10": "µg/m³", "PM2.5": "µg/m³"}

# ============================ DATA + STATISTICS =============================
df = pd.read_excel(DATA, sheet_name="Data")
d = df[df.STATION == "CEN"].copy()
for c in CAND + ["NO2"]:
    d[c] = pd.to_numeric(d[c], errors="coerce")
n_no2 = int(d["NO2"].notna().sum())


def pearson_slide(x, y):
    # computational formula from Session 6
    n = len(x)
    num = (x * y).sum() - x.sum() * y.sum() / n
    den = np.sqrt((x**2).sum() - x.sum()**2 / n) * np.sqrt((y**2).sum() - y.sum()**2 / n)
    return num / den


rows = []
for c in CAND:
    s = d[["NO2", c]].dropna()
    x, y = s[c].values, s["NO2"].values
    cov = ((x - x.mean()) * (y - y.mean())).sum() / len(x)
    r = pearson_slide(x, y)
    b1, b0 = np.polyfit(x, y, 1)
    rows.append({"var": c, "r": r, "cov": cov, "r2": r**2, "n": len(x), "b0": b0, "b1": b1})
stats = pd.DataFrame(rows).sort_values("r", key=lambda v: v.abs(), ascending=False).reset_index(drop=True)
top = stats.iloc[0]
indep = stats[~stats["var"].isin(["NOX", "NO"])].iloc[0]
st = stats.set_index("var")


def r2_score(y, yh):
    return 1 - ((y - yh)**2).sum() / ((y - y.mean())**2).sum()


LIN_VARS = ["NOX", "CO", "O3"]
fits = {}
for c in LIN_VARS:
    s = d[["NO2", c]].dropna(); s = s[s[c] > 0]
    x, y = s[c].values, s["NO2"].values
    p_lin = np.polyfit(x, y, 1); p_log = np.polyfit(np.log(x), y, 1); p_q = np.polyfit(x, y, 2)
    fits[c] = {"x": x, "y": y, "p_lin": p_lin, "p_log": p_log, "p_q": p_q,
               "lin": r2_score(y, np.polyval(p_lin, x)),
               "log": r2_score(y, np.polyval(p_log, np.log(x))),
               "cuad": r2_score(y, np.polyval(p_q, x))}

s_log = d[["NO2", "CO"]].dropna()
p90 = s_log["NO2"].quantile(0.90)
y_bin = (s_log["NO2"] > p90).astype(int)
logit = sm.Logit(y_bin, sm.add_constant(s_log["CO"])).fit(disp=0)
lb0, lb1 = logit.params["const"], logit.params["CO"]
co50 = -lb0 / lb1
p_at1 = 1 / (1 + np.exp(-(lb0 + lb1 * 1.0)))
n_alto = int(y_bin.sum())


def f(v, dd=3):
    return f"{v:.{dd}f}"


def pct(v):
    return f"{v * 100:.0f} %"


# ============================ FIGURES =======================================
def fig_formulas():
    p = OUT + r"\formulas_cov_r.png"
    fig = plt.figure(figsize=(6.0, 2.3))
    fig.text(0.5, 0.78, r"$cov(X,Y)=\frac{\sum (x_i-\bar{x})(y_i-\bar{y})}{n}$",
             ha="center", va="center", fontsize=15)
    fig.text(0.5, 0.30,
             r"$r_{X,Y}=\frac{\sum x_i y_i-\frac{\sum x_i\sum y_i}{n}}"
             r"{\sqrt{\sum x_i^2-\frac{(\sum x_i)^2}{n}}\;\sqrt{\sum y_i^2-\frac{(\sum y_i)^2}{n}}}$",
             ha="center", va="center", fontsize=15)
    plt.savefig(p, dpi=200, bbox_inches="tight"); plt.close()
    return p


def fig_scatter():
    p = OUT + r"\dispersion_regresion_no2.png"
    fig, ax = plt.subplots(2, 3, figsize=(10.5, 6.6))
    for axx, c in zip(ax.ravel(), stats["var"]):
        s = d[["NO2", c]].dropna(); row = st.loc[c]
        axx.scatter(s[c], s["NO2"], s=5, alpha=0.18, color=NAVY_HEX, edgecolors="none")
        xs = np.linspace(s[c].min(), s[c].max(), 50)
        axx.plot(xs, row["b0"] + row["b1"] * xs, color=AMBER_HEX, linewidth=2)
        is_top = c == top["var"]
        axx.set_title(f"NO₂ vs {SHORT[c]}   (r = {row['r']:+.3f})",
                      color=(AMBER_HEX if is_top else NAVY_HEX), weight="bold", fontsize=10)
        axx.set_xlabel(f"{SHORT[c]} ({UNIT[c]})", fontsize=8); axx.set_ylabel("NO₂ (ppm)", fontsize=8)
        axx.grid(linestyle=":", alpha=0.4)
        if is_top:
            for sp in axx.spines.values():
                sp.set_edgecolor(AMBER_HEX); sp.set_linewidth(2)
    plt.tight_layout(); plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


def fig_linearity():
    p = OUT + r"\linealidad_no2.png"
    fig, ax = plt.subplots(1, 3, figsize=(11, 3.9))
    for axx, c in zip(ax, LIN_VARS):
        F = fits[c]; xs = np.linspace(F["x"].min(), F["x"].max(), 200)
        axx.scatter(F["x"], F["y"], s=4, alpha=0.18, color="#9A9A9A", edgecolors="none")
        axx.plot(xs, np.polyval(F["p_lin"], xs), color="#1F3FD0", lw=2,
                 label=f"Lineal  (R² = {F['lin']:.3f})")
        axx.plot(xs, np.polyval(F["p_log"], np.log(xs)), color="#2E9E3E", lw=2,
                 label=f"Logarítmico  (R² = {F['log']:.3f})")
        axx.plot(xs, np.polyval(F["p_q"], xs), color="#D02020", lw=2,
                 label=f"Cuadrático  (R² = {F['cuad']:.3f})")
        axx.set_title(f"NO₂ vs {SHORT[c]}", color=NAVY_HEX, weight="bold", fontsize=11)
        axx.set_xlabel(f"{SHORT[c]} ({UNIT[c]})", fontsize=9); axx.set_ylabel("NO₂ (ppm)", fontsize=9)
        axx.legend(fontsize=7.5, loc="upper right" if c == "O3" else "lower right")
        axx.grid(linestyle=":", alpha=0.4)
    plt.tight_layout(); plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


def fig_logistic():
    p = OUT + r"\logistica_no2_alto_co.png"
    fig, ax = plt.subplots(figsize=(7.4, 4.3))
    x = s_log["CO"].values; yb = y_bin.values
    ax.scatter(x[yb == 1], yb[yb == 1], s=10, color="#3CB043", alpha=0.5, label="NO₂ alto (> P90)")
    ax.scatter(x[yb == 0], yb[yb == 0], s=10, color="#D9534F", alpha=0.35, label="NO₂ no alto")
    xs = np.linspace(x.min(), x.max(), 300)
    ax.plot(xs, 1 / (1 + np.exp(-(lb0 + lb1 * xs))), color="#1F3FD0", lw=2.2,
            label="P(NO₂ alto | CO) = 1 / (1 + e^-(b₀ + b₁·CO))")
    ax.axhline(0.5, color="#777777", ls="--", lw=1)
    ax.axvline(co50, color="#777777", ls="--", lw=1)
    ax.annotate(f"P = 0.5 con CO = {co50:.2f} ppm", xy=(co50, 0.5), xytext=(co50 + 0.6, 0.38),
                fontsize=8.5, color="#333333", arrowprops=dict(arrowstyle="->", color="#777777"))
    ax.set_title("Variable continua (CO) frente a variable discreta (NO₂ alto)",
                 color=NAVY_HEX, weight="bold", fontsize=11)
    ax.set_xlabel("CO (ppm)"); ax.set_ylabel("Probabilidad / clase")
    ax.legend(fontsize=8, loc="center right"); ax.grid(linestyle=":", alpha=0.4)
    plt.tight_layout(); plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


# ============================ DOCX HELPERS ==================================
def set_cell_bg(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"); tcPr.append(shd)


def add_run(p, t, bold=False, italic=False, color=None, size=None, font="Calibri"):
    r = p.add_run(t); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    if size: r.font.size = Pt(size)
    r.font.name = font; return r


def add_para(doc, t="", bold=False, italic=False, color=None, size=10,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, sb=0, sa=5, keep=False):
    p = doc.add_paragraph(); p.alignment = align
    p.paragraph_format.space_before = Pt(sb); p.paragraph_format.space_after = Pt(sa)
    p.paragraph_format.keep_with_next = keep
    if t: add_run(p, t, bold, italic, color, size)
    return p


def add_heading(doc, t, level=1):
    return add_para(doc, t, bold=True, color=NAVY, size=(13 if level == 1 else 11),
                    align=WD_ALIGN_PARAGRAPH.LEFT, sb=(10 if level == 1 else 7), sa=3, keep=True)


def add_caption(doc, t):
    return add_para(doc, t, italic=True, color=RGBColor(0x40, 0x40, 0x40), size=9,
                    align=WD_ALIGN_PARAGRAPH.CENTER, sb=2, sa=8)


def add_pic(doc, path, w):
    doc.add_picture(path, width=Inches(w))
    pp = doc.paragraphs[-1]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.keep_with_next = True   # keep figure with its caption


def make_table(doc, headers, rows_, caption, widths, hl=None):
    tbl = doc.add_table(rows=1 + len(rows_), cols=len(headers))
    tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, w in enumerate(widths):
        for row in tbl.rows:
            row.cells[i].width = Inches(w)
    for c in tbl.rows[0].cells: set_cell_bg(c, NAVY)
    for ci, h in enumerate(headers):
        pp = tbl.rows[0].cells[ci].paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(pp, h, bold=True, color=WHITE, size=9)
    for ri, rd in enumerate(rows_):
        dr = tbl.rows[ri + 1]
        bg = HL if hl == ri else (WHITE if ri % 2 == 0 else GRAY_L)
        for c in dr.cells: set_cell_bg(c, bg)
        for ci, v in enumerate(rd):
            pp = dr.cells[ci].paragraphs[0]
            pp.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
            add_run(pp, str(v), size=9, bold=(hl == ri))
    add_caption(doc, caption)


def add_apa(doc, t):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.5); p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.space_after = Pt(6); add_run(p, t, size=10)


def clear_body_after_table1(doc):
    body = doc.element.body; sect = body.find(qn("w:sectPr")); c = 0; cut = None
    for i, ch in enumerate(list(body)):
        if ch.tag == qn("w:tbl"):
            c += 1
            if c == 2: cut = i; break
    if cut is not None:
        for el in [e for e in list(body)[cut + 1:] if e.tag != qn("w:sectPr")]: body.remove(el)
    if sect is not None and sect not in list(body): body.append(sect)


# ============================ BUILD =========================================
def build():
    os.makedirs(OUT, exist_ok=True)
    p_form, p_sc, p_lin, p_lg = fig_formulas(), fig_scatter(), fig_linearity(), fig_logistic()
    Fx, Fc = fits["NOX"], fits["CO"]

    doc = Document(TEMPLATE)
    hdr = doc.tables[0].rows[0].cells[1]
    hdr.paragraphs[2].runs[0].text = "Maestría en Ciencias de los Datos | 3er Semestre"
    hdr.paragraphs[3].runs[0].text = "Materia: Visualización de Grandes Bases de Datos"
    tp = doc.paragraphs[1]; tp.clear()
    r = tp.add_run("Tarea 1.4. Correlación lineal entre el NO₂ y otros contaminantes\n"
                   "Estación de monitoreo Centro, Zona Metropolitana de Guadalajara, 2024")
    r.bold = True; r.font.size = Pt(13); r.font.color.rgb = NAVY; r.font.name = "Calibri"
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def si(idx, text):
        cell = doc.tables[1].rows[idx].cells[1]; cell.paragraphs[0].runs[0].text = text
        for extra in cell.paragraphs[0].runs[1:]: extra.text = ""
    si(1, "24 de septiembre de 2026")
    si(2, "Tarea 1.4. Correlación con NO₂")
    si(3, "Python 3.13 (pandas, NumPy, Matplotlib, statsmodels)")
    si(4, "SIMAJ, base de calidad del aire 2024 (BD_2024)")
    si(5, "Equipo personal, entorno conda MCD")
    si(6, "Calidad del aire ZMG 2024, estación Centro (CEN)")
    clear_body_after_table1(doc)

    # ---------------- 1. Introduccion
    add_para(doc, sb=8, sa=0)
    add_heading(doc, "1. Introducción")
    add_para(doc,
        "El dióxido de nitrógeno (NO₂) es uno de los contaminantes criterio que se asocian con "
        "mayor claridad a las emisiones del tránsito vehicular. En esta actividad se busca "
        "identificar cuál de los contaminantes registrados por la red de monitoreo guarda la "
        "relación lineal más fuerte con el NO₂, a partir del coeficiente de correlación de "
        "Pearson, y sustentar el resultado mediante gráficas de dispersión y rectas de regresión. "
        "Siguiendo los temas de la sesión 6, el análisis se complementa con una revisión de la "
        "linealidad de las relaciones y con un modelo para una variable de respuesta discreta. "
        "Los datos provienen de la base horaria de calidad del aire de 2024 del Sistema de "
        "Monitoreo Atmosférico de Jalisco (SIMAJ).")
    add_heading(doc, "1.1 Estación de monitoreo", level=2)
    add_para(doc,
        "En las tareas 1.0 a 1.2 se trabajó con la estación Country, por ser la más próxima a mi "
        "domicilio. Al revisar la base, sin embargo, se comprobó que esta estación no reporta "
        "mediciones de NO₂ durante 2024, lo que impide realizar el análisis solicitado. Por ese "
        f"motivo se utilizó la estación Centro, ya empleada en la Tarea 1.3, que cuenta con "
        f"{n_no2:,} registros horarios de NO₂ y mide además O₃, NO, NOx, CO, PM10 y PM2.5. El SO₂ "
        "no se registra en esta estación, por lo que quedó fuera del análisis.")

    # ---------------- 2. Metodologia
    add_heading(doc, "2. Metodología")
    add_heading(doc, "2.1 Covarianza y coeficiente de correlación", level=2)
    add_para(doc,
        "La covarianza mide el grado en que dos variables varían de manera conjunta respecto a "
        "sus medias. Su signo indica el sentido de la relación, pero su magnitud depende de las "
        "unidades de medición, por lo que no permite comparar pares de variables expresadas en "
        "escalas distintas. El coeficiente de correlación de Pearson resuelve esta limitación al "
        "estandarizar la covarianza por la dispersión de ambas variables. Su valor se ubica entre "
        "−1 y 1: los extremos corresponden a una relación lineal perfecta, negativa o positiva, y "
        "los valores cercanos a cero indican ausencia de relación lineal. Ambos estadísticos se "
        "calcularon con las expresiones vistas en clase (Figura 1), considerando para cada par "
        "únicamente las horas en que las dos variables tienen lectura válida.", keep=True)
    add_pic(doc, p_form, 4.4)
    add_caption(doc, "Figura 1. Expresiones de la covarianza y del coeficiente de correlación lineal de Pearson.")

    add_heading(doc, "2.2 Recta de regresión y coeficiente de determinación", level=2)
    add_para(doc,
        "Para cada contaminante X se ajustó por mínimos cuadrados ordinarios la recta "
        "NO₂ = b₀ + b₁X. El coeficiente de determinación (R²) indica la proporción de la "
        "variabilidad del NO₂ que reproduce el modelo; en la regresión lineal simple coincide con "
        "el cuadrado de r.")

    add_heading(doc, "2.3 Linealidad de la relación", level=2)
    add_para(doc,
        "Un coeficiente de correlación moderado no implica necesariamente que la asociación sea "
        "débil, ya que la relación puede existir sin seguir una línea recta. Para valorar este "
        "aspecto, en el NOx, el CO y el O₃ se compararon tres especificaciones con base en su R²: "
        "lineal (ŷ = b₀ + b₁x), logarítmica (ŷ = b₀ + b₁·log x) y cuadrática "
        "(ŷ = b₀ + b₁x + b₂x²).")

    add_heading(doc, "2.4 Variable continua y variable discreta", level=2)
    add_para(doc,
        "Por último, el NO₂ se transformó en una variable binaria que distingue las horas de "
        f"concentración alta, definidas como aquellas que superan el percentil 90 ({f(p90, 3)} ppm), "
        "del resto de las observaciones. La relación entre esta variable y el CO se modeló "
        "mediante una regresión logística, P(Y = 1 | x) = 1 / (1 + e^−(b₀ + b₁x)), que estima la "
        "probabilidad de una hora con NO₂ alto en función de la concentración de CO.")

    # ---------------- 3. Resultados
    add_heading(doc, "3. Resultados")
    add_heading(doc, "3.1 Correlación lineal con el NO₂", level=2)
    add_para(doc,
        "La Tabla 1 resume los estadísticos de cada par, ordenados por el valor absoluto de r. "
        "Las covarianzas confirman el sentido de cada relación, pero sus magnitudes no son "
        f"comparables entre sí. La del PM10 ({f(st.loc['PM10','cov'], 4)}) es varios cientos de "
        f"veces mayor que la del NOx ({f(st.loc['NOX','cov'], 6)}) únicamente porque las "
        "partículas se expresan en µg/m³ y presentan mayor variabilidad absoluta. Por esa razón, "
        "la comparación entre contaminantes se basa en el coeficiente de correlación.", keep=True)
    trows = [[NAMES[rr["var"]], f"{rr['r']:+.3f}", f"{rr['cov']:.6f}", f(rr["r2"]), f"{int(rr['n']):,}"]
             for _, rr in stats.iterrows()]
    make_table(doc, ["Contaminante", "r de Pearson", "Covarianza", "R²", "Pares (n)"], trows,
               "Tabla 1. Covarianza, correlación lineal y R² de cada contaminante con el NO₂, "
               "estación Centro, 2024.", [2.9, 1.0, 1.1, 0.8, 0.9], hl=0)

    add_para(doc,
        "La Figura 2 presenta las gráficas de dispersión con su recta de regresión y el valor de "
        "r en el encabezado de cada panel.", sb=4, sa=3, keep=True)
    add_pic(doc, p_sc, 6.4)
    add_caption(doc, "Figura 2. Gráficas de dispersión y rectas de regresión del NO₂ frente a cada "
                     "contaminante, estación Centro, 2024. El panel con borde ámbar corresponde a la "
                     "correlación más alta.")
    add_para(doc,
        "Con excepción del O₃, todos los contaminantes presentan correlación positiva: la nube de "
        "puntos asciende hacia la derecha y la pendiente de la recta es positiva. El O₃ muestra el "
        f"patrón inverso. Los coeficientes más altos corresponden al NOx (r = {f(st.loc['NOX','r'])}) "
        f"y al CO (r = {f(st.loc['CO','r'])}), cuyas nubes son también las más compactas alrededor "
        "de la recta. En el PM10, el NO y el PM2.5 la dispersión es mayor, lo que concuerda con "
        "correlaciones de intensidad moderada.")

    add_heading(doc, "3.2 ¿Relación lineal o no lineal?", level=2)
    add_para(doc,
        "En los paneles del NOx y del CO, la recta se ajusta bien en las concentraciones bajas, "
        "pero sobrestima el NO₂ en los valores altos, donde los puntos quedan sistemáticamente por "
        "debajo de ella. Este patrón sugiere una relación cóncava, por lo que se compararon las "
        "tres especificaciones descritas en la metodología (Tabla 2 y Figura 3).", keep=True)
    lrows = [[NAMES[c], f(fits[c]["lin"]), f(fits[c]["log"]), f(fits[c]["cuad"])] for c in LIN_VARS]
    make_table(doc, ["Contaminante", "R² lineal", "R² logarítmico", "R² cuadrático"], lrows,
               "Tabla 2. Comparación del ajuste de los modelos lineal, logarítmico y cuadrático "
               "para el NO₂.", [2.6, 1.3, 1.4, 1.4])
    add_pic(doc, p_lin, 6.5)
    add_caption(doc, "Figura 3. Ajustes lineal, logarítmico y cuadrático del NO₂ frente al NOx, el "
                     "CO y el O₃, estación Centro, 2024.")
    add_para(doc,
        "En los tres casos el modelo logarítmico obtiene el mayor R². Para el NOx, la proporción "
        f"de variabilidad explicada pasa de {f(Fx['lin'])} con la recta a {f(Fx['log'])}, y para "
        f"el CO, de {f(Fc['lin'])} a {f(Fc['log'])}. La curva logarítmica sigue la forma de la nube "
        "con mucha mayor fidelidad que la recta, lo que indica que el coeficiente de Pearson, al "
        "medir solo la componente lineal, subestima la intensidad de la asociación. El NO₂ crece "
        "con rapidez cuando el NOx o el CO aumentan desde niveles bajos y lo hace de forma más "
        "gradual en concentraciones altas. Este comportamiento es compatible con la química de la "
        "atmósfera urbana: las emisiones vehiculares recientes están compuestas sobre todo por NO, "
        "cuya oxidación a NO₂ depende de la disponibilidad de ozono y otros oxidantes, de manera "
        "que en los episodios de mayor emisión la proporción de NO₂ dentro del NOx disminuye. El "
        "modelo cuadrático también mejora el ajuste, aunque en menor medida, y su forma de "
        "parábola hace que la curva descienda dentro del propio rango observado, algo que carece "
        "de interpretación física para estos contaminantes.")

    add_heading(doc, "3.3 Probabilidad de horas con NO₂ alto", level=2)
    add_para(doc,
        f"De las {len(y_bin):,} horas con lectura de ambos contaminantes, {n_alto} superan el "
        "percentil 90 del NO₂. La Figura 4 muestra la regresión logística ajustada entre esta "
        "variable binaria y el CO.", keep=True)
    add_pic(doc, p_lg, 5.4)
    add_caption(doc, "Figura 4. Regresión logística de la probabilidad de una hora con NO₂ alto en "
                     "función de la concentración de CO, estación Centro, 2024.")
    add_para(doc,
        f"La probabilidad estimada es inferior a {pct(p_at1)} cuando el CO no rebasa 1 ppm, crece "
        f"con rapidez a partir de ese nivel y alcanza 50 % cuando el CO llega a {f(co50, 2)} ppm. "
        f"El modelo obtiene un pseudo R² de McFadden de {f(logit.prsquared)}, valor que indica un "
        "buen ajuste para este tipo de regresión. El resultado complementa el análisis de "
        "correlación: además de variar en la misma dirección, las concentraciones altas de CO "
        "permiten anticipar con buena probabilidad las horas de mayor concentración de NO₂.")

    # ---------------- 4. Conclusiones
    add_heading(doc, "4. Conclusiones")
    add_para(doc,
        "El contaminante con el mayor coeficiente de correlación lineal con el NO₂ en la estación "
        f"Centro es el NOx (r = {f(top['r'])}). Este resultado debe interpretarse con reserva, ya "
        "que el NOx se calcula como la suma de NO y NO₂ y, por lo tanto, contiene a la variable con "
        "la que se compara. Una parte de la correlación obedece, entonces, a la forma en que está "
        "definido el indicador.")
    add_para(doc,
        "Entre los contaminantes que no incluyen al NO₂ en su cálculo, el CO presenta la "
        f"correlación más alta (r = {f(indep['r'])}), prácticamente igual a la del NOx. La afinidad "
        "entre ambos se explica porque comparten su principal fuente, la combustión en vehículos "
        "automotores, y responden de la misma forma a las condiciones de dispersión de la "
        f"atmósfera. El O₃, en cambio, se relaciona de manera inversa con el NO₂ (r = "
        f"{f(st.loc['O3','r'])}), debido a que su formación fotoquímica ocurre en las horas de "
        "mayor radiación solar, cuando el NO₂ desciende.")
    add_para(doc,
        "El análisis de linealidad matiza estos resultados. Las relaciones del NO₂ con el NOx y con "
        f"el CO son cóncavas, y un modelo logarítmico explica cerca de {pct(Fx['log'])} y "
        f"{pct(Fc['log'])} de la variabilidad del NO₂, frente a {pct(Fx['lin'])} y "
        f"{pct(Fc['lin'])} de la recta. El coeficiente de Pearson resulta adecuado para jerarquizar "
        "a los contaminantes, pero subestima la fuerza de la asociación. La regresión logística, "
        "por su parte, confirma que el CO es un buen indicador de las horas con NO₂ alto. En todos "
        "los casos se trata de asociaciones estadísticas, que por sí mismas no establecen "
        "relaciones de causalidad.")

    add_heading(doc, "Referencias")
    add_apa(doc, "Hosmer, D. W., Lemeshow, S., & Sturdivant, R. X. (2013). Applied logistic "
                 "regression (3.ª ed.). Wiley.")
    add_apa(doc, "Ruiz Chávez, C. F. (2026). Visualización de grandes bases de datos: Sesión 6 "
                 "[Diapositivas de clase]. Maestría en Ciencia de los Datos, Centro Universitario de "
                 "Ciencias Económico Administrativas, Universidad de Guadalajara.")
    add_apa(doc, "Secretaría de Medio Ambiente y Desarrollo Territorial. (2024). Sistema de "
                 "Monitoreo Atmosférico de Jalisco (SIMAJ): base de datos de calidad del aire 2024. "
                 "Gobierno del Estado de Jalisco. https://siga.jalisco.gob.mx/aire/")
    add_apa(doc, "Seinfeld, J. H., & Pandis, S. N. (2016). Atmospheric chemistry and physics: From "
                 "air pollution to climate change (3.ª ed.). Wiley.")
    out = OUT + r"\Tarea1_4_VisualizacionGBD_CarlosPulidoRosas.docx"
    doc.save(out); print("DOCX:", out); return out


if __name__ == "__main__":
    build()
