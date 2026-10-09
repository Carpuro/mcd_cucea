# -*- coding: utf-8 -*-
"""
Tarea 1.5 - Adjustments to the NO2 correlation of Tarea 1.4 (station Centro), per Session 7:
homogeneous corpus (same hours for every pollutant), exclusion of Tukey outliers before the
regressions, regression lines through the origin (0, 0), heatmaps and scatter + boxplot combos.
"""
import os
import subprocess
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.gridspec as gridspec
import statsmodels.api as sm
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"C:\Users\carlo\GitHub\mcd_cucea"
TEMPLATE = BASE + r"\_recursos\plantillas\plantilla_informe_MCD.docx"
COURSE = BASE + r"\3er-semestre\visualizacion-grandes-bases-datos"
DATA = COURSE + r"\datos\BD_2024_calidad_aire.xlsx"
OUT = COURSE + r"\tareas\Tarea1_5"
SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"
NAVY = RGBColor(0x1F, 0x49, 0x7D); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY_L = RGBColor(0xF2, 0xF2, 0xF2); HL = RGBColor(0xFC, 0xE4, 0xB8)
NAVY_HEX, AMBER_HEX, RED_HEX = "#1F497D", "#D48C00", "#C0392B"

CAND = ["O3", "NO", "NOX", "CO", "PM10", "PM2.5"]
ALL = ["NO2"] + CAND
NAMES = {"NO2": "NO₂ (dióxido de nitrógeno)", "O3": "O₃ (ozono)", "NO": "NO (monóxido de nitrógeno)",
         "NOX": "NOx (óxidos de nitrógeno)", "CO": "CO (monóxido de carbono)",
         "PM10": "PM10 (partículas < 10 µm)", "PM2.5": "PM2.5 (partículas < 2.5 µm)"}
SHORT = {"NO2": "NO₂", "O3": "O₃", "NO": "NO", "NOX": "NOx", "CO": "CO", "PM10": "PM10", "PM2.5": "PM2.5"}
UNIT = {"NO2": "ppm", "O3": "ppm", "NO": "ppm", "NOX": "ppm", "CO": "ppm", "PM10": "µg/m³", "PM2.5": "µg/m³"}
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre",
         "octubre", "noviembre", "diciembre"]

# ============================ DATA ==========================================
df = pd.read_excel(DATA, sheet_name="Data")
d = df[df.STATION == "CEN"].copy()
for c in ALL:
    d[c] = pd.to_numeric(d[c], errors="coerce")
d["DATE"] = pd.to_datetime(d["DATE"])
n_total = len(d)
n_valid = {c: int(d[c].notna().sum()) for c in ALL}
no2_months = sorted(d.loc[d["NO2"].notna(), "DATE"].dt.month.unique())

# Tarea 1.4 baseline: every pair used its own set of hours
r14 = {c: d[["NO2", c]].dropna().corr().iloc[0, 1] for c in CAND}
n14 = {c: len(d[["NO2", c]].dropna()) for c in CAND}

# 1) homogeneous corpus: hours with a valid reading of all seven variables
H = d.dropna(subset=ALL).copy()
n_h = len(H)

# 2) Tukey outliers (1.5 IQR) per variable; an hour is excluded if it is atypical in any
#    variable, so every regression keeps using exactly the same hours
fence = {}
mask = pd.Series(True, index=H.index)
for c in ALL:
    q1, q3 = np.percentile(H[c], [25, 75]); iqr = q3 - q1
    lo, hi = q1 - 1.5 * iqr, q3 + 1.5 * iqr
    out_c = (H[c] < lo) | (H[c] > hi)
    fence[c] = {"q1": q1, "q3": q3, "iqr": iqr, "lo": lo, "hi": hi, "n_out": int(out_c.sum())}
    mask &= ~out_c
H["atipico"] = ~mask
C = H[mask].copy()
n_c = len(C)

# 3) correlation at each stage + regression through the origin on the clean corpus
rows = []
for c in CAND:
    x, y = C[c].values, C["NO2"].values
    o = sm.OLS(y, x).fit()                       # no constant: NO2 = b * X
    ic = sm.OLS(y, sm.add_constant(x)).fit()     # reference fit with intercept
    sse = ((y - o.params[0] * x) ** 2).sum()
    rows.append({"var": c, "r14": r14[c], "n14": n14[c], "rh": H[c].corr(H["NO2"]),
                 "rc": C[c].corr(C["NO2"]), "b": o.params[0], "r2_0": o.rsquared,
                 "r2_media": 1 - sse / ((y - y.mean()) ** 2).sum(),
                 "a_ic": ic.params[0], "b_ic": ic.params[1], "r2_ic": ic.rsquared})
stats = pd.DataFrame(rows).sort_values("rc", key=lambda v: v.abs(), ascending=False).reset_index(drop=True)
st = stats.set_index("var")
top = stats.iloc[0]
indep = stats[~stats["var"].isin(["NOX", "NO"])].iloc[0]
R_C = C[ALL].corr()
R_14 = pd.DataFrame(index=ALL, columns=ALL, dtype=float)
for a in ALL:
    for b_ in ALL:
        R_14.loc[a, b_] = d[[a, b_]].dropna().corr().iloc[0, 1]


def f(v, dd=3):
    return f"{v:.{dd}f}"


def pct(v):
    return f"{v * 100:.1f} %"


def lista(vs):
    vs = [SHORT[v] for v in vs]
    return vs[0] if len(vs) == 1 else ", el ".join(vs[:-1]) + " y el " + vs[-1]


def num(v):
    return f"{v:,}"


# ============================ FIGURES =======================================
def fig_formulas():
    p = OUT + r"\formulas_origen.png"
    fig = plt.figure(figsize=(5.0, 1.6))
    fig.text(0.5, 0.78, r"$\hat{y}=b\,x,\qquad b=\frac{\sum x_i\,y_i}{\sum x_i^{2}}$",
             ha="center", va="center", fontsize=15)
    fig.text(0.5, 0.26, r"$R^{2}_{0}=1-\frac{\sum (y_i-b\,x_i)^{2}}{\sum y_i^{2}}$",
             ha="center", va="center", fontsize=15)
    plt.savefig(p, dpi=200, bbox_inches="tight"); plt.close()
    return p


def fig_boxplots():
    p = OUT + r"\boxplots_atipicos.png"
    fig, ax = plt.subplots(1, 7, figsize=(11, 3.4))
    for axx, c in zip(ax, ALL):
        axx.boxplot(H[c], widths=0.55, patch_artist=True,
                    boxprops=dict(facecolor="#DCE6F1", edgecolor=NAVY_HEX),
                    medianprops=dict(color=AMBER_HEX, lw=2), whiskerprops=dict(color=NAVY_HEX),
                    capprops=dict(color=NAVY_HEX),
                    flierprops=dict(marker="o", ms=2.5, mfc=RED_HEX, mec="none", alpha=0.5))
        axx.axhline(fence[c]["hi"], color=RED_HEX, ls="--", lw=1)
        axx.set_title(SHORT[c], color=NAVY_HEX, weight="bold", fontsize=10)
        axx.set_xticks([]); axx.set_ylabel(UNIT[c], fontsize=8); axx.tick_params(labelsize=7)
        axx.set_xlabel(f"{fence[c]['n_out']} atípicos", fontsize=8, color=RED_HEX)
        axx.grid(axis="y", linestyle=":", alpha=0.4)
    plt.tight_layout(); plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


def fig_heatmaps():
    p = OUT + r"\mapas_calor_correlacion.png"
    fig, ax = plt.subplots(1, 2, figsize=(11, 4.9))
    labs = [SHORT[c] for c in ALL]
    for axx, M, t in zip(ax, [R_14, R_C], ["Tarea 1.4: cada par con sus propias horas",
                                           "Tarea 1.5: corpus homogéneo sin atípicos"]):
        im = axx.imshow(M.values.astype(float), cmap="coolwarm", vmin=-1, vmax=1)
        axx.set_xticks(range(7)); axx.set_yticks(range(7))
        axx.set_xticklabels(labs, fontsize=9); axx.set_yticklabels(labs, fontsize=9)
        for i in range(7):
            for j in range(7):
                v = float(M.values[i, j])
                axx.text(j, i, f"{v:.2f}", ha="center", va="center", fontsize=8,
                         color="white" if abs(v) > 0.6 else "black")
        axx.set_title(t, color=NAVY_HEX, weight="bold", fontsize=11)
    fig.colorbar(im, ax=list(ax), shrink=0.8, label="r de Pearson")
    plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


def fig_scatter_origin():
    p = OUT + r"\dispersion_origen_no2.png"
    fig, ax = plt.subplots(2, 3, figsize=(10.5, 6.8))
    for axx, c in zip(ax.ravel(), stats["var"]):
        row = st.loc[c]; x, y = C[c].values, C["NO2"].values
        axx.scatter(x, y, s=5, alpha=0.22, color=NAVY_HEX, edgecolors="none")
        xs = np.linspace(0, x.max(), 50)
        axx.plot(xs, row["b"] * xs, color=AMBER_HEX, lw=2.2, label="Recta por el origen")
        axx.plot(xs, row["a_ic"] + row["b_ic"] * xs, color="#777777", lw=1.2, ls="--",
                 label="Recta con intercepto")
        axx.set_xlim(left=0); axx.set_ylim(bottom=0)
        is_top = c == top["var"]
        axx.set_title(f"NO₂ vs {SHORT[c]}   (r = {row['rc']:+.3f})",
                      color=(AMBER_HEX if is_top else NAVY_HEX), weight="bold", fontsize=10)
        axx.set_xlabel(f"{SHORT[c]} ({UNIT[c]})", fontsize=8); axx.set_ylabel("NO₂ (ppm)", fontsize=8)
        axx.text(0.03, 0.95, f"b = {row['b']:.4g}\nR²₀ = {row['r2_0']:.3f}", transform=axx.transAxes,
                 va="top", fontsize=8, bbox=dict(fc="white", ec="#BBBBBB", alpha=0.85))
        axx.grid(linestyle=":", alpha=0.4)
        if is_top:
            for sp in axx.spines.values():
                sp.set_edgecolor(AMBER_HEX); sp.set_linewidth(2)
    ax[0, 0].legend(fontsize=7, loc="lower right")
    plt.tight_layout(); plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


def fig_marginal():
    # scatter + distribution (Session 7): boxplots on the margins, excluded hours in red
    p = OUT + r"\dispersion_distribucion_nox_co.png"
    fig = plt.figure(figsize=(11, 5.2))
    outer = gridspec.GridSpec(1, 2, wspace=0.22)
    kept, out_ = H[~H["atipico"]], H[H["atipico"]]
    for k, c in enumerate(["NOX", "CO"]):
        g = gridspec.GridSpecFromSubplotSpec(2, 2, subplot_spec=outer[k], width_ratios=[1, 6],
                                             height_ratios=[6, 1], wspace=0.05, hspace=0.05)
        axm = fig.add_subplot(g[0, 1]); axl = fig.add_subplot(g[0, 0], sharey=axm)
        axb = fig.add_subplot(g[1, 1], sharex=axm)
        axm.scatter(out_[c], out_["NO2"], s=6, color=RED_HEX, alpha=0.35, edgecolors="none",
                    label=f"Horas excluidas ({len(out_):,})")
        axm.scatter(kept[c], kept["NO2"], s=6, color="#7FB3D5", alpha=0.6, edgecolors="none",
                    label=f"Horas conservadas ({len(kept):,})")
        xs = np.linspace(0, kept[c].max(), 50)
        axm.plot(xs, st.loc[c, "b"] * xs, color=NAVY_HEX, lw=2.2, label="Recta por el origen")
        axm.set_xlim(left=0); axm.set_ylim(bottom=0)
        axm.set_title(f"NO₂ vs {SHORT[c]}  (r sin atípicos = {st.loc[c, 'rc']:.3f})",
                      color=NAVY_HEX, weight="bold", fontsize=10)
        axm.grid(linestyle=":", alpha=0.5); axm.legend(fontsize=7, loc="upper left")
        plt.setp(axm.get_xticklabels(), visible=False); plt.setp(axm.get_yticklabels(), visible=False)
        bp = dict(patch_artist=True, boxprops=dict(facecolor="white", edgecolor="#333333"),
                  medianprops=dict(color="#333333"), flierprops=dict(marker="o", ms=2, mfc="none"))
        axl.boxplot(H["NO2"], widths=0.6, **bp); axl.set_xticks([])
        axl.set_ylabel("NO₂ (ppm)", fontsize=8); axl.tick_params(labelsize=7)
        axb.boxplot(H[c], vert=False, widths=0.6, **bp); axb.set_yticks([])
        axb.set_xlabel(f"{SHORT[c]} ({UNIT[c]})", fontsize=8); axb.tick_params(labelsize=7)
    plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


# ============================ DOCX HELPERS ==================================
def set_cell_bg(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"); tcPr.append(shd)


def add_run(p, t, bold=False, italic=False, color=None, size=None, font="Calibri"):
    r = p.add_run(t); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    if size: r.font.size = Pt(size)
    r.font.name = font; return r


def add_para(doc, t="", bold=False, italic=False, color=None, size=10,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, sb=0, sa=5, keep=False):
    p = doc.add_paragraph(); p.alignment = align
    p.paragraph_format.space_before = Pt(sb); p.paragraph_format.space_after = Pt(sa)
    p.paragraph_format.keep_with_next = keep
    if t: add_run(p, t, bold, italic, color, size)
    return p


def add_heading(doc, t, level=1):
    return add_para(doc, t, bold=True, color=NAVY, size=(13 if level == 1 else 11),
                    align=WD_ALIGN_PARAGRAPH.LEFT, sb=(10 if level == 1 else 7), sa=3, keep=True)


def add_caption(doc, t):
    return add_para(doc, t, italic=True, color=RGBColor(0x40, 0x40, 0x40), size=9,
                    align=WD_ALIGN_PARAGRAPH.CENTER, sb=2, sa=8)


def add_pic(doc, path, w):
    doc.add_picture(path, width=Inches(w))
    pp = doc.paragraphs[-1]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.keep_with_next = True   # keep figure with its caption


def make_table(doc, headers, rows_, caption, widths, hl=None):
    tbl = doc.add_table(rows=1 + len(rows_), cols=len(headers))
    tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, w in enumerate(widths):
        for row in tbl.rows:
            row.cells[i].width = Inches(w)
    for c in tbl.rows[0].cells: set_cell_bg(c, NAVY)
    for ci, h in enumerate(headers):
        pp = tbl.rows[0].cells[ci].paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(pp, h, bold=True, color=WHITE, size=9)
    for ri, rd in enumerate(rows_):
        dr = tbl.rows[ri + 1]
        bg = HL if hl == ri else (WHITE if ri % 2 == 0 else GRAY_L)
        for c in dr.cells: set_cell_bg(c, bg)
        for ci, v in enumerate(rd):
            pp = dr.cells[ci].paragraphs[0]
            pp.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
            add_run(pp, str(v), size=9, bold=(hl == ri))
            pp.paragraph_format.keep_with_next = True   # keep the table and its caption together
    for c in tbl.rows[0].cells: c.paragraphs[0].paragraph_format.keep_with_next = True
    add_caption(doc, caption)


def add_apa(doc, t):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.5); p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.space_after = Pt(6); add_run(p, t, size=10)


def clear_body_after_table1(doc):
    body = doc.element.body; sect = body.find(qn("w:sectPr")); c = 0; cut = None
    for i, ch in enumerate(list(body)):
        if ch.tag == qn("w:tbl"):
            c += 1
            if c == 2: cut = i; break
    if cut is not None:
        for el in [e for e in list(body)[cut + 1:] if e.tag != qn("w:sectPr")]: body.remove(el)
    if sect is not None and sect not in list(body): body.append(sect)


# ============================ BUILD =========================================
def build():
    os.makedirs(OUT, exist_ok=True)
    p_form, p_box, p_hm = fig_formulas(), fig_boxplots(), fig_heatmaps()
    p_sc, p_mg = fig_scatter_origin(), fig_marginal()
    meses = f"{MESES[no2_months[0] - 1]} a {MESES[no2_months[-1] - 1]}"

    doc = Document(TEMPLATE)
    hdr = doc.tables[0].rows[0].cells[1]
    hdr.paragraphs[2].runs[0].text = "Maestría en Ciencias de los Datos | 3er Semestre"
    hdr.paragraphs[3].runs[0].text = "Materia: Visualización de Grandes Bases de Datos"
    tp = doc.paragraphs[1]; tp.clear()
    r = tp.add_run("Tarea 1.5. Ajustes a la correlación lineal con el NO₂\n"
                   "Estación de monitoreo Centro, Zona Metropolitana de Guadalajara, 2024")
    r.bold = True; r.font.size = Pt(13); r.font.color.rgb = NAVY; r.font.name = "Calibri"
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def si(idx, text):
        cell = doc.tables[1].rows[idx].cells[1]; cell.paragraphs[0].runs[0].text = text
        for extra in cell.paragraphs[0].runs[1:]: extra.text = ""
    si(1, "8 de octubre de 2026")
    si(2, "Tarea 1.5. Ajustes a la correlación con NO₂")
    si(3, "Python 3.13 (pandas, NumPy, Matplotlib, statsmodels)")
    si(4, "SIMAJ, base de calidad del aire 2024 (BD_2024)")
    si(5, "Equipo personal, entorno conda MCD")
    si(6, "Calidad del aire ZMG 2024, estación Centro (CEN)")
    clear_body_after_table1(doc)

    # ---------------- 1. Introduccion
    add_para(doc, sb=8, sa=0)
    add_heading(doc, "1. Introducción")
    add_para(doc,
        "En la Tarea 1.4 busqué qué contaminante tiene la correlación lineal más alta con el "
        "dióxido de nitrógeno (NO₂) en la estación Centro del SIMAJ. Después de la sesión 7 quedó "
        "claro que a ese análisis le faltaban tres ajustes, y en esta tarea los hago. El primero "
        "es trabajar con un corpus homogéneo, es decir, que todos los contaminantes se comparen "
        "usando exactamente las mismas horas. El segundo es quitar los valores atípicos antes de "
        "ajustar las regresiones. El tercero es que las rectas de regresión pasen por el origen "
        "(0, 0). Con los datos ya ajustados vuelvo a buscar el contaminante con mayor correlación y "
        "reviso el resultado con mapas de calor, gráficas de dispersión, diagramas de caja y "
        "rectas de regresión.")

    # ---------------- 2. Metodologia
    add_heading(doc, "2. Metodología")
    add_heading(doc, "2.1 Corpus homogéneo", level=2)
    add_para(doc,
        "En la Tarea 1.4 calculé cada correlación con las horas en que el NO₂ y el otro "
        "contaminante tenían dato. El problema es que cada par terminó usando un número distinto "
        f"de horas (entre {num(min(n14.values()))} y {num(max(n14.values()))}), y no siempre las "
        "mismas. Así, si un coeficiente salía más alto que otro, no se podía saber si era por el "
        "contaminante o porque se estaban comparando horas diferentes. A eso se suma que la "
        f"estación solo midió NO₂ de {meses} de 2024, mientras que el O₃, el CO y las partículas "
        "se midieron casi todo el año. Para que la comparación fuera justa, me quedé solo con las "
        "horas que tienen dato de las siete variables al mismo tiempo (Tabla 1).")
    make_table(doc, ["Etapa", "Horas", "% del total"],
               [["Registros horarios de la estación Centro, 2024", num(n_total), "100.0 %"],
                [f"Con lectura de NO₂ ({meses})", num(n_valid['NO2']), pct(n_valid['NO2'] / n_total)],
                ["Corpus homogéneo (lectura de las siete variables)", num(n_h), pct(n_h / n_total)],
                ["Corpus homogéneo sin valores atípicos", num(n_c), pct(n_c / n_total)]],
               "Tabla 1. Construcción del corpus de análisis, estación Centro, 2024.",
               [3.9, 0.9, 1.0], hl=3)

    add_heading(doc, "2.2 Exclusión de valores atípicos", level=2)
    add_para(doc,
        "Para encontrar los atípicos usé el criterio de Tukey de las tareas 1.0 y 1.1. Un valor es "
        "atípico si queda fuera del intervalo [Q1 − 1.5·RIC, Q3 + 1.5·RIC], donde RIC = Q3 − Q1. "
        "Calculé las cercas de cada variable sobre el corpus homogéneo. Si una hora era atípica en "
        "cualquiera de las siete variables, la quité completa, para que todas las regresiones "
        "siguieran usando las mismas horas y no se perdiera lo ganado en el paso anterior. Como "
        "todas las distribuciones tienen la cola hacia la derecha, la cerca inferior siempre salió "
        "negativa, así que solo se quitaron valores altos (Tabla 2 y Figura 1).", keep=True)
    frows = [[NAMES[c], f(fence[c]["q1"], 3 if UNIT[c] == "ppm" else 1),
              f(fence[c]["q3"], 3 if UNIT[c] == "ppm" else 1),
              f(fence[c]["hi"], 4 if UNIT[c] == "ppm" else 1), fence[c]["n_out"]] for c in ALL]
    make_table(doc, ["Variable", "Q1", "Q3", "Cerca superior", "Horas atípicas"], frows,
               f"Tabla 2. Cercas de Tukey y horas atípicas por variable, corpus homogéneo "
               f"(n = {num(n_h)}).", [2.6, 0.8, 0.8, 1.1, 1.0])
    add_pic(doc, p_box, 6.5)
    add_caption(doc, "Figura 1. Diagramas de caja de cada variable en el corpus homogéneo. La línea "
                     "punteada roja marca la cerca superior de Tukey; los puntos rojos son los "
                     "valores atípicos.")
    n_excl = n_h - n_c
    add_para(doc,
        f"En total quité {num(n_excl)} horas, el {pct(n_excl / n_h)} del corpus. El NO es el que "
        f"tiene más atípicos ({fence['NO']['n_out']}). Tiene sentido, porque sale directo del escape "
        "de los coches y se dispara en las horas de más tráfico. Muchas veces una misma hora es "
        "atípica en varios contaminantes a la vez, y por eso el total de horas quitadas es menor "
        "que la suma de la última columna de la Tabla 2.")

    add_heading(doc, "2.3 Recta de regresión que cruza por el origen", level=2)
    add_para(doc,
        "En la Tarea 1.4 usé la recta NO₂ = b₀ + b₁X. Ahora ajusté por mínimos cuadrados el modelo "
        "sin intercepto NO₂ = b·X, que obliga a la recta a pasar por (0, 0). La idea es que, si no "
        "hay emisiones, tampoco debería haber NO₂. Las fórmulas de la pendiente y del coeficiente "
        "de determinación están en la Figura 2.", keep=True)
    add_pic(doc, p_form, 2.9)
    add_caption(doc, "Figura 2. Pendiente y coeficiente de determinación de la recta de regresión "
                     "que cruza por el origen.")
    add_para(doc,
        "Hay que tener cuidado con el R²₀. En un modelo sin intercepto el R² se mide contra cero y "
        "no contra el promedio, y por eso suele salir más alto de lo que en realidad ajusta el "
        "modelo. No se puede comparar con el R² de una recta normal ni con r². Por esa razón usé "
        "el coeficiente de Pearson para ordenar los contaminantes, y la recta por el origen solo "
        "para ver la proporción entre el NO₂ y cada contaminante.")

    # ---------------- 3. Resultados
    add_heading(doc, "3. Resultados")
    add_heading(doc, "3.1 Efecto de los ajustes sobre el coeficiente de correlación", level=2)
    add_para(doc,
        "En la Tabla 3 se ve cómo cambia la correlación de cada contaminante con el NO₂ después de "
        "cada ajuste. La Figura 3 compara las matrices de correlación de la Tarea 1.4 y de esta "
        "tarea en mapas de calor, donde el rojo intenso indica correlación positiva fuerte y el "
        "azul, negativa.", keep=True)
    crow = [[NAMES[rr["var"]], f"{rr['r14']:+.3f}", f"{rr['rh']:+.3f}", f"{rr['rc']:+.3f}",
             f"{rr['rc'] - rr['r14']:+.3f}"] for _, rr in stats.iterrows()]
    make_table(doc, ["Contaminante", "r Tarea 1.4", "r corpus homogéneo", "r sin atípicos",
                     "Cambio"], crow,
               "Tabla 3. Coeficiente de correlación de Pearson con el NO₂ en cada etapa del ajuste, "
               "ordenado por el valor absoluto de r sin atípicos.", [2.5, 0.9, 1.3, 1.1, 0.8], hl=0)
    add_pic(doc, p_hm, 6.5)
    add_caption(doc, "Figura 3. Mapas de calor de la matriz de correlación de Pearson antes "
                     "(Tarea 1.4) y después de los ajustes (Tarea 1.5).")
    sub_ = stats[stats["rc"] < stats["r14"]]
    add_para(doc,
        "Usar el corpus homogéneo casi no movió los coeficientes (ninguno cambió más de 0.02). O "
        "sea, que en la Tarea 1.4 cada par tuviera distintas horas no estaba distorsionando mucho "
        "los resultados. Quitar los atípicos, en cambio, sí hizo una diferencia, y no en la misma "
        f"dirección para todos. En el NOx la correlación subió de {f(st.loc['NOX','r14'])} a "
        f"{f(st.loc['NOX','rc'])}, y en el CO de {f(st.loc['CO','r14'])} a {f(st.loc['CO','rc'])}.")
    add_para(doc,
        "Para entender por qué, hay que recordar lo que se vio en la Tarea 1.4. Con el NOx y el CO "
        "la nube de puntos se curva: cuando sus concentraciones son muy altas, el NO₂ ya no sube al "
        "mismo ritmo y los puntos quedan por debajo de la recta. Esos puntos son justo los que el "
        "criterio de Tukey marca como atípicos. Al quitarlos, lo que queda es casi una línea recta, "
        "y r sube. En el " + lista(sub_["var"]) + " pasó lo contrario y la correlación bajó. Ahí "
        "unas cuantas horas con valores extremos en las dos variables a la vez jalaban la recta "
        "hacia arriba, porque los puntos muy alejados del resto pesan mucho en el cálculo "
        "(puntos de alta palanca). Sin ellos, la relación resulta más débil de lo que parecía.")

    add_heading(doc, "3.2 Gráficas de dispersión y rectas por el origen", level=2)
    add_para(doc,
        "La Figura 4 muestra las gráficas de dispersión ya sin atípicos, con la recta que pasa por "
        "el origen en ámbar y, para comparar, la recta con intercepto en gris punteado. Los "
        "valores de los dos modelos están en la Tabla 4.")
    add_pic(doc, p_sc, 6.4)
    add_caption(doc, f"Figura 4. Dispersión del NO₂ frente a cada contaminante con rectas de "
                     f"regresión que cruzan por el origen, corpus homogéneo sin atípicos "
                     f"(n = {num(n_c)}). El panel con borde ámbar corresponde a la correlación más alta.")
    orow = [[NAMES[rr["var"]], f"{rr['rc']:+.3f}", f"{rr['b']:.4g}", f(rr["r2_0"]),
             f"{rr['a_ic']:+.4f}", f"{rr['b_ic']:.4g}", f(rr["r2_ic"])] for _, rr in stats.iterrows()]
    make_table(doc, ["Contaminante", "r", "b (origen)", "R²₀", "b₀ (intercepto)", "b₁", "R²"], orow,
               "Tabla 4. Regresión por el origen y regresión con intercepto del NO₂ sobre cada "
               "contaminante, corpus homogéneo sin atípicos.", [2.3, 0.6, 0.8, 0.6, 1.0, 0.7, 0.6], hl=0)
    add_para(doc,
        "En el NOx y el CO las dos rectas casi coinciden. Cuando se deja libre el intercepto, sale "
        f"prácticamente en cero ({st.loc['NOX','a_ic']:+.4f} y {st.loc['CO','a_ic']:+.4f} ppm), así "
        "que obligar a la recta a pasar por el origen no le quita casi nada al ajuste. Las "
        "pendientes también se pueden leer de forma sencilla. La del NOx "
        f"(b = {f(st.loc['NOX','b'])}) quiere decir que el NO₂ es más o menos el "
        f"{pct(st.loc['NOX','b'])} del NOx medido. La del CO (b = {f(st.loc['CO','b'], 4)}) indica "
        f"que por cada ppm de CO hay en promedio {f(st.loc['CO','b'] * 1000, 1)} ppb de NO₂.")
    add_para(doc,
        "Con los demás contaminantes las dos rectas se separan bastante. El caso más claro es el "
        f"O₃. Su correlación con el NO₂ es negativa (r = {f(st.loc['O3','rc'])}), pero la recta "
        f"forzada al origen sale con pendiente positiva (b = {f(st.loc['O3','b'])}), porque si "
        "tiene que arrancar en (0, 0), la única manera de acercarse a los puntos es subiendo. Para "
        "el O₃ esta restricción no sirve, ya que contradice lo que muestran los datos, y tiene "
        "lógica, porque el O₃ no sale de la misma fuente que el NO₂ sino que se forma consumiéndolo. "
        "Algo parecido pasa con el R²₀ de estas variables, que sale alto aunque el ajuste sea malo. "
        "Si se mide contra el promedio, la recta por el origen ajusta peor que una línea "
        "horizontal en el " + lista(stats.loc[stats["r2_media"] < 0, "var"]) + ". Por eso no usé "
        "el R²₀ para comparar.")

    add_heading(doc, "3.3 Dispersión y distribución del NOx y el CO", level=2)
    add_para(doc,
        "En la Figura 5 junté la gráfica de dispersión con los diagramas de caja de cada eje, como "
        "en el ejemplo de la sesión 7, para los dos contaminantes con mayor correlación. Las horas "
        "que quité por atípicas están en rojo.", keep=True)
    add_pic(doc, p_mg, 6.6)
    add_caption(doc, "Figura 5. Dispersión del NO₂ frente al NOx y al CO con diagramas de caja "
                     "marginales, corpus homogéneo. En rojo, las horas excluidas por atípicas; la "
                     "recta se ajustó solo con las horas conservadas.")
    add_para(doc,
        "Los diagramas de caja muestran que el NO₂, el NOx y el CO tienen una cola larga hacia los "
        "valores altos. Las horas en rojo están en esa cola, en la zona donde la nube se "
        "abre y se dobla. Las horas que se quedaron (en azul) forman una banda angosta y casi "
        "recta que sale del origen, sobre todo en el NOx. Con esta gráfica se entiende por qué la "
        "correlación subió al quitar los atípicos y por qué la recta por el origen describe bien la "
        "relación en los niveles normales de contaminación.")

    # ---------------- 4. Conclusiones
    add_heading(doc, "4. Conclusiones")
    add_para(doc,
        "Después de los tres ajustes, el contaminante con mayor correlación lineal con el NO₂ sigue "
        f"siendo el NOx, que pasó de r = {f(top['r14'])} en la Tarea 1.4 a r = {f(top['rc'])}. Como "
        "ya había mencionado, este resultado hay que tomarlo con reserva, porque el NOx se calcula "
        "sumando el NO y el NO₂. Es decir, ya trae dentro la variable con la que se está comparando.")
    add_para(doc,
        "Si solo se consideran los contaminantes que no incluyen al NO₂, el de mayor correlación es "
        f"el CO (r = {f(indep['rc'])}). Su relación con el NO₂ también mejoró con los ajustes y la "
        f"describe bien una recta por el origen (b = {f(indep['b'], 4)}; con intercepto libre, este "
        f"sale en {indep['a_ic']:+.4f} ppm). Los dos vienen sobre todo de los coches, por eso crecen "
        f"de forma proporcional. El O₃ mantiene su relación inversa (r = {f(st.loc['O3','rc'])}) y, "
        "junto con el NO y las partículas, queda bastante por debajo del NOx y el CO.")
    add_para(doc,
        "Lo que más me llamó la atención es que los atípicos estaban afectando los resultados de "
        "la Tarea 1.4 en sentidos opuestos. En el NOx y el CO escondían parte de la relación "
        "lineal, y en el NO y las partículas la hacían ver más fuerte de lo que es. También quedó "
        "claro que la recta por el origen solo funciona cuando el intercepto libre ya sale cerca "
        "de cero, como con el NOx y el CO. Con el O₃ da una pendiente con el signo al revés. Una "
        "correlación alta tampoco significa que un contaminante cause al otro; aquí lo que tienen "
        "en común es la fuente.")

    add_heading(doc, "Referencias")
    add_apa(doc, "Montgomery, D. C., Peck, E. A., & Vining, G. G. (2021). Introduction to linear "
                 "regression analysis (6.ª ed.). Wiley.")
    add_apa(doc, "Ruiz Chávez, C. F. (2026). Visualización de grandes bases de datos: Sesión 7 "
                 "[Diapositivas de clase]. Maestría en Ciencia de los Datos, Centro Universitario de "
                 "Ciencias Económico Administrativas, Universidad de Guadalajara.")
    add_apa(doc, "Secretaría de Medio Ambiente y Desarrollo Territorial. (2024). Sistema de "
                 "Monitoreo Atmosférico de Jalisco (SIMAJ): base de datos de calidad del aire 2024. "
                 "Gobierno del Estado de Jalisco. https://siga.jalisco.gob.mx/aire/")
    add_apa(doc, "Seinfeld, J. H., & Pandis, S. N. (2016). Atmospheric chemistry and physics: From "
                 "air pollution to climate change (3.ª ed.). Wiley.")
    add_apa(doc, "Tukey, J. W. (1977). Exploratory data analysis. Addison-Wesley.")
    out = OUT + r"\Tarea1_5_VisualizacionGBD_CarlosPulidoRosas.docx"
    doc.save(out); print("DOCX:", out)
    subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", OUT, out],
                   check=True, capture_output=True)
    print("PDF:", out.replace(".docx", ".pdf"))
    return out


if __name__ == "__main__":
    build()
