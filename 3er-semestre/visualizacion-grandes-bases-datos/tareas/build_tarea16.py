# -*- coding: utf-8 -*-
"""
Tarea 1.6 - NO2 imputation with the simple linear models of tasks 1.4 and 1.5, then exploration
of NO2 vs PM10 / PM2.5 with scatter matrices and multiple linear regression (Session 8:
multidimensionality, collinearity, linear and quadratic regression planes). Station Centro.
"""
import os
import subprocess
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import statsmodels.api as sm
from statsmodels.stats.outliers_influence import variance_inflation_factor
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"C:\Users\carlo\GitHub\mcd_cucea"
TEMPLATE = BASE + r"\_recursos\plantillas\plantilla_informe_MCD.docx"
COURSE = BASE + r"\3er-semestre\visualizacion-grandes-bases-datos"
DATA = COURSE + r"\datos\BD_2024_calidad_aire.xlsx"
OUT = COURSE + r"\tareas\Tarea1_6"
SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"
NAVY = RGBColor(0x1F, 0x49, 0x7D); WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAY_L = RGBColor(0xF2, 0xF2, 0xF2); HL = RGBColor(0xFC, 0xE4, 0xB8)
NAVY_HEX, AMBER_HEX, RED_HEX, GREEN_HEX = "#1F497D", "#D48C00", "#C0392B", "#2E8B57"

ALL = ["NO2", "O3", "NO", "NOX", "CO", "PM10", "PM2.5"]
SHORT = {"NO2": "NO₂", "CO": "CO", "PM10": "PM10", "PM2.5": "PM2.5"}
UNIT = {"NO2": "ppm", "CO": "ppm", "PM10": "µg/m³", "PM2.5": "µg/m³"}
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre",
         "octubre", "noviembre", "diciembre"]
MES3 = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]


def tukey(s):
    q1, q3 = np.percentile(s, [25, 75]); iqr = q3 - q1
    return q1 - 1.5 * iqr, q3 + 1.5 * iqr


# ============================ DATA ==========================================
df = pd.read_excel(DATA, sheet_name="Data")
d = df[df.STATION == "CEN"].copy()
for c in ALL:
    d[c] = pd.to_numeric(d[c], errors="coerce")
d["DATE"] = pd.to_datetime(d["DATE"]); d["mes"] = d["DATE"].dt.month
n_total = len(d)
n_obs = int(d["NO2"].notna().sum())
n_falt = n_total - n_obs
n_falt_co = int((d["NO2"].isna() & d["CO"].notna()).sum())
n_falt_sin = n_falt - n_falt_co
n_falt_sd = int((d["NO2"].isna() & (d["mes"] >= 9)).sum())

# ---- the two simple models of the previous tasks (NO2 explained by CO, the best predictor that
#      does not contain NO2 and the only one with year-round coverage)
p14 = d[["NO2", "CO"]].dropna()
b1_14, b0_14 = np.polyfit(p14["CO"], p14["NO2"], 1)            # Tarea 1.4: intercept, all hours
H = d.dropna(subset=ALL)
keep = pd.Series(True, index=H.index)
for c in ALL:
    lo, hi = tukey(H[c]); keep &= (H[c] >= lo) & (H[c] <= hi)
C15 = H[keep]
b_15 = (C15["CO"] * C15["NO2"]).sum() / (C15["CO"] ** 2).sum()  # Tarea 1.5: through the origin
co_lim = tukey(H["CO"])[1]                                       # upper Tukey fence of CO in 1.5
co_max15 = C15["CO"].max()


def pred14(co):
    return b0_14 + b1_14 * co


def pred15(co):
    return b_15 * co


def pred_comb(co):
    # model 1.5 inside the range it was fitted on, model 1.4 beyond the CO fence
    co = np.asarray(co, dtype=float)
    return np.where(co <= co_lim, pred15(co), pred14(co))


def pred_acot(co):
    # model 1.5 limited to the range it was fitted on: beyond the CO fence the estimate is held
    # at its value on the fence instead of extrapolating the line (NO2 flattens at high CO)
    co = np.asarray(co, dtype=float)
    return pred15(np.minimum(co, co_lim))


def metrics(y, yh):
    e = y - yh
    return {"mae": np.abs(e).mean(), "rmse": np.sqrt((e ** 2).mean()), "sesgo": e.mean()}


# in-sample error on the observed hours, overall and by CO range
y_o, co_o = p14["NO2"].values, p14["CO"].values
alto = co_o > co_lim
ev = {}
for k, fn in [("1.4", pred14), ("1.5", pred15), ("comb", pred_comb), ("acot", pred_acot)]:
    yh = fn(co_o)
    ev[k] = {"todo": metrics(y_o, yh), "bajo": metrics(y_o[~alto], yh[~alto]),
             "alto": metrics(y_o[alto], yh[alto])}
n_alto_obs = int(alto.sum())

# out-of-sample check: refit both models on Sep-Oct and predict Nov-Dec
tr, te = p14[d.loc[p14.index, "mes"] <= 10], p14[d.loc[p14.index, "mes"] >= 11]
Htr = H[H["mes"] <= 10]; ktr = pd.Series(True, index=Htr.index)
for c in ALL:
    lo, hi = tukey(Htr[c]); ktr &= (Htr[c] >= lo) & (Htr[c] <= hi)
Ctr = Htr[ktr]
bt15 = (Ctr["CO"] * Ctr["NO2"]).sum() / (Ctr["CO"] ** 2).sum()
bt1, bt0 = np.polyfit(tr["CO"], tr["NO2"], 1)
lim_tr = tukey(Htr["CO"])[1]
yt, cot = te["NO2"].values, te["CO"].values
oos = {"1.4": metrics(yt, bt0 + bt1 * cot), "1.5": metrics(yt, bt15 * cot),
       "comb": metrics(yt, np.where(cot <= lim_tr, bt15 * cot, bt0 + bt1 * cot)),
       "acot": metrics(yt, bt15 * np.minimum(cot, lim_tr))}

# imputation
d["NO2_c"] = d["NO2"]
m_imp = d["NO2"].isna() & d["CO"].notna()
d.loc[m_imp, "NO2_c"] = pred_acot(d.loc[m_imp, "CO"])
d["origen"] = np.where(d["NO2"].notna(), "observado", np.where(m_imp, "imputado", "faltante"))
n_imp_alto = int((m_imp & (d["CO"] > co_lim)).sum())
mens = d.groupby("mes").agg(obs=("NO2", "mean"), comp=("NO2_c", "mean"),
                            n_obs=("NO2", "count"), n_imp=("origen", lambda s: (s == "imputado").sum()))

# ---- exploration NO2 vs PM10 / PM2.5: complete series, hours with the three variables,
#      Tukey outliers excluded as in Tarea 1.5; the observed-only corpus is kept for comparison
VARS = ["NO2", "PM10", "PM2.5"]


def corpus(frame, ycol):
    z = frame.dropna(subset=[ycol, "PM10", "PM2.5"])[[ycol, "PM10", "PM2.5", "origen", "mes"]]
    z = z.rename(columns={ycol: "NO2"})
    k = pd.Series(True, index=z.index)
    for c in VARS:
        lo, hi = tukey(z[c]); k &= (z[c] >= lo) & (z[c] <= hi)
    return z, z[k]


Zc_all, Zc = corpus(d, "NO2_c")
Zo_all, Zo = corpus(d[d["NO2"].notna()], "NO2")


def modelos(Z):
    y = Z["NO2"]; res = {}
    for k, cols in [("PM10", ["PM10"]), ("PM2.5", ["PM2.5"]), ("lineal", ["PM10", "PM2.5"])]:
        res[k] = sm.OLS(y, sm.add_constant(Z[cols])).fit()
    Q = Z.assign(PM10_2=Z["PM10"] ** 2)
    res["cuad"] = sm.OLS(y, sm.add_constant(Q[["PM10", "PM10_2", "PM2.5"]])).fit()
    X = sm.add_constant(Z[["PM10", "PM2.5"]]).values
    res["vif"] = variance_inflation_factor(X, 1)
    res["r"] = Z[VARS].corr()
    return res


Mc, Mo = modelos(Zc), modelos(Zo)

# descriptive figures of the particles over the whole year (all valid hours) and daily means
# against the 24-h WHO (2021) air quality guidelines: 45 ug/m3 PM10, 15 ug/m3 PM2.5
OMS = {"PM10": 45, "PM2.5": 15}
pm = {}
for c in ["PM10", "PM2.5"]:
    dia = d.set_index("DATE")[c].resample("D").mean().dropna()
    pm[c] = {"media": d[c].mean(), "mediana": d[c].median(), "dias": len(dia),
             "exc": int((dia > OMS[c]).sum())}
ratio_pm = (d["PM2.5"] / d["PM10"]).replace(np.inf, np.nan).median()

# NO2 coverage in the whole SIMAJ network: first month with readings at any station
red = df.assign(NO2=pd.to_numeric(df["NO2"], errors="coerce"), mes=pd.to_datetime(df["DATE"]).dt.month)
red_cov = red.dropna(subset=["NO2"]).groupby("STATION")["mes"].min()
n_est = red["STATION"].nunique(); n_est_no2 = len(red_cov); mes_ini_red = int(red_cov.min())


# coarse fraction: PM10 = PM2.5 + (PM10 - PM2.5); same fit as the PM10 + PM2.5 model, but the two
# regressors no longer share the fine particles, so collinearity drops and each one is interpretable
def gruesa(Z):
    G = Z.assign(PMg=Z["PM10"] - Z["PM2.5"])
    X = sm.add_constant(G[["PM2.5", "PMg"]])
    return {"m": sm.OLS(G["NO2"], X).fit(), "vif": variance_inflation_factor(X.values, 1),
            "r": G[["PM2.5", "PMg"]].corr().iloc[0, 1], "neg": int((G["PMg"] < 0).sum())}


Gc, Go = gruesa(Zc), gruesa(Zo)

# daily cycle with observed NO2 only (September to December), each variable relative to its mean
d["hora"] = d["DATE"].dt.hour
hor = d[d["NO2"].notna()].groupby("hora")[["NO2", "PM10", "PM2.5", "CO"]].mean()
hor_rel = hor / hor.mean()
no2_med_o = Zo["NO2"].mean()


def f(v, dd=3):
    return f"{v:.{dd}f}"


def pct(v):
    return f"{v * 100:.1f} %"


def num(v):
    return f"{int(v):,}"


# ============================ FIGURES =======================================
def fig_imputacion():
    p = OUT + r"\modelos_imputacion.png"
    fig, ax = plt.subplots(figsize=(8.2, 4.4))
    ax.scatter(co_o[~alto], y_o[~alto], s=5, alpha=0.3, color="#7FB3D5", edgecolors="none",
               label="Horas observadas")
    ax.scatter(co_o[alto], y_o[alto], s=6, alpha=0.4, color=RED_HEX, edgecolors="none",
               label=f"Horas observadas con CO > {co_lim:.2f} ppm")
    xs = np.linspace(0, co_o.max(), 300)
    ax.plot(xs, pred14(xs), color="#777777", lw=1.5, ls="--",
            label=f"Modelo 1.4: NO₂ = {b0_14:.4f} + {b1_14:.4f}·CO")
    ax.plot(xs, pred15(xs), color=AMBER_HEX, lw=1.5, ls=":",
            label=f"Modelo 1.5: NO₂ = {b_15:.4f}·CO")
    ax.plot(xs, pred_comb(xs), color="#8E44AD", lw=1.2, label="Regla combinada (1.5 y 1.4)")
    ax.plot(xs, pred_acot(xs), color=NAVY_HEX, lw=2.6, label="Modelo 1.5 acotado (imputación)")
    ax.axvline(co_lim, color=RED_HEX, lw=1, ls="--")
    ax.text(co_lim + 0.08, 0.002, "cerca de Tukey del CO", color=RED_HEX, fontsize=8)
    ax.set_ylim(0, 0.12); ax.set_xlim(0, co_o.max() * 1.02)
    ax.set_xlabel("CO (ppm)"); ax.set_ylabel("NO₂ (ppm)")
    ax.set_title("Modelos de regresión simple usados para imputar el NO₂", color=NAVY_HEX,
                 weight="bold", fontsize=11)
    ax.legend(fontsize=7.5, loc="upper left"); ax.grid(linestyle=":", alpha=0.4)
    plt.tight_layout(); plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


def fig_serie():
    p = OUT + r"\serie_no2_imputada.png"
    dd = d.set_index("DATE")
    obs = dd["NO2"].resample("D").mean()
    imp = dd["NO2_c"].where(dd["origen"] == "imputado").resample("D").mean()
    fig, ax = plt.subplots(figsize=(10, 3.6))
    ax.plot(imp.index, imp.values, color=AMBER_HEX, lw=1.1, label="NO₂ imputado (promedio diario)")
    ax.plot(obs.index, obs.values, color=NAVY_HEX, lw=1.1, label="NO₂ observado (promedio diario)")
    ax.set_ylabel("NO₂ (ppm)"); ax.grid(linestyle=":", alpha=0.4); ax.legend(fontsize=8, loc="upper right")
    ax.set_title("Serie diaria de NO₂ completada, estación Centro, 2024", color=NAVY_HEX,
                 weight="bold", fontsize=11)
    ax.set_xticks(pd.date_range("2024-01-01", "2024-12-01", freq="MS"))
    ax.set_xticklabels(MES3, fontsize=8)
    plt.tight_layout(); plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


def fig_matriz():
    # Session 8: Y vs each X, and X vs X to look for collinearity
    p = OUT + r"\matriz_dispersion_no2_pm.png"
    pares = [("PM10", "NO2"), ("PM2.5", "NO2"), ("PM2.5", "PM10")]
    fig, ax = plt.subplots(1, 3, figsize=(11, 3.9))
    for axx, (x, y) in zip(ax, pares):
        Z = Zc; col = np.where(Z["origen"] == "imputado", AMBER_HEX, NAVY_HEX)
        axx.scatter(Z[x], Z[y], s=4, alpha=0.25, c=col, edgecolors="none")
        b1, b0 = np.polyfit(Z[x], Z[y], 1); xs = np.linspace(0, Z[x].max(), 50)
        axx.plot(xs, b0 + b1 * xs, color=RED_HEX, lw=2)
        r = Z[[x, y]].corr().iloc[0, 1]
        tag = "  (colinealidad)" if y != "NO2" else ""
        axx.set_title(f"{SHORT[y]} vs {SHORT[x]}{tag}", color=NAVY_HEX, weight="bold", fontsize=10)
        axx.set_xlabel(f"{SHORT[x]} ({UNIT[x]})\nr = {r:.3f}", fontsize=8)
        axx.set_ylabel(f"{SHORT[y]} ({UNIT[y]})", fontsize=8); axx.grid(linestyle=":", alpha=0.4)
        axx.set_xlim(left=0); axx.set_ylim(bottom=0)
    ax[0].scatter([], [], color=NAVY_HEX, s=12, label="NO₂ observado")
    ax[0].scatter([], [], color=AMBER_HEX, s=12, label="NO₂ imputado")
    ax[0].legend(fontsize=7, loc="upper left")
    plt.tight_layout(); plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


def fig_planos():
    # Session 8: linear plane and quadratic surface in 3D
    p = OUT + r"\planos_regresion_3d.png"
    Z = Zc; samp = Z.sample(min(1500, len(Z)), random_state=1)
    g1 = np.linspace(0, Z["PM10"].max(), 25); g2 = np.linspace(0, Z["PM2.5"].max(), 25)
    G1, G2 = np.meshgrid(g1, g2)
    fig = plt.figure(figsize=(11, 4.9))
    for k, (key, title) in enumerate([("lineal", "Modelo lineal: NO₂ = b₀ + b₁·PM10 + b₂·PM2.5"),
                                      ("cuad", "Modelo cuadrático: NO₂ = b₀ + b₁·PM10 + b₂·PM10² + b₃·PM2.5")]):
        ax = fig.add_subplot(1, 2, k + 1, projection="3d")
        pr = Mc[key].params
        if key == "lineal":
            S = pr["const"] + pr["PM10"] * G1 + pr["PM2.5"] * G2
        else:
            S = pr["const"] + pr["PM10"] * G1 + pr["PM10_2"] * G1 ** 2 + pr["PM2.5"] * G2
        ax.plot_surface(G1, G2, S, color="#5B6BD6", alpha=0.45, edgecolor="#333366", lw=0.2)
        ax.scatter(samp["PM10"], samp["PM2.5"], samp["NO2"], s=3, color="#4A7FA7", alpha=0.45)
        ax.set_xlabel("PM10 (µg/m³)", fontsize=8); ax.set_ylabel("PM2.5 (µg/m³)", fontsize=8)
        ax.set_zlabel("NO₂ (ppm)", fontsize=8); ax.tick_params(labelsize=6)
        ax.set_title(f"{title}\nR² = {Mc[key].rsquared:.3f}", color=NAVY_HEX, weight="bold", fontsize=9)
        ax.view_init(elev=22, azim=-58)
    plt.tight_layout(); plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


def fig_horario():
    p = OUT + r"\ciclo_diario.png"
    fig, ax = plt.subplots(figsize=(9, 3.8))
    sty = {"NO2": (NAVY_HEX, "-", 2.4), "PM10": (AMBER_HEX, "-", 2.0), "PM2.5": (RED_HEX, "-", 2.0),
           "CO": ("#777777", "--", 1.4)}
    for c, (col, ls, lw) in sty.items():
        ax.plot(hor_rel.index, hor_rel[c], color=col, ls=ls, lw=lw, marker="o", ms=3, label=SHORT[c])
    ax.axhline(1, color="#333333", lw=0.8)
    ax.axvspan(10.5, 17.5, color="#F2F2F2", zorder=0)
    ax.set_ylim(0, hor_rel.values.max() * 1.25)
    ax.text(14, hor_rel.values.max() * 1.18, "mediodía y tarde", ha="center", va="top", fontsize=8,
            color="#555555")
    ax.set_xticks(range(0, 24, 2)); ax.set_xlabel("Hora del día"); ax.set_ylabel("Concentración / promedio")
    ax.set_title("Ciclo diario promedio, horas con NO₂ observado (septiembre a diciembre de 2024)",
                 color=NAVY_HEX, weight="bold", fontsize=10)
    ax.legend(fontsize=8, ncol=4, loc="upper left"); ax.grid(linestyle=":", alpha=0.4)
    plt.tight_layout(); plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


def fig_coef():
    # partial effect of PM2.5: sign in the simple model vs the multiple model, both corpora
    p = OUT + r"\coeficientes_pm25.png"
    fig, ax = plt.subplots(figsize=(7.2, 3.3))
    labs = ["Simple\n(solo PM2.5)", "Múltiple\n(PM10 + PM2.5)"]
    xs = np.arange(2); w = 0.36
    for j, (M, nm, colr) in enumerate([(Mo, "Solo horas observadas", NAVY_HEX),
                                       (Mc, "Serie completa (con imputados)", AMBER_HEX)]):
        vals = [M["PM2.5"].params["PM2.5"], M["lineal"].params["PM2.5"]]
        lo = [M["PM2.5"].conf_int().loc["PM2.5"], M["lineal"].conf_int().loc["PM2.5"]]
        err = np.array([[v - l[0] for v, l in zip(vals, lo)], [l[1] - v for v, l in zip(vals, lo)]])
        ax.bar(xs + (j - 0.5) * w, np.array(vals) * 1000, w, yerr=err * 1000, color=colr,
               capsize=3, label=nm)
    ax.axhline(0, color="#333333", lw=0.8); ax.set_xticks(xs); ax.set_xticklabels(labs, fontsize=9)
    ax.set_ylabel("Coeficiente de PM2.5\n(ppb de NO₂ por µg/m³)", fontsize=8)
    ax.set_title("Coeficiente del PM2.5 con intervalo de confianza de 95 %", color=NAVY_HEX,
                 weight="bold", fontsize=10)
    ax.legend(fontsize=7.5); ax.grid(axis="y", linestyle=":", alpha=0.4)
    plt.tight_layout(); plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


# ============================ DOCX HELPERS ==================================
def set_cell_bg(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"); tcPr.append(shd)


def add_run(p, t, bold=False, italic=False, color=None, size=None, font="Calibri"):
    r = p.add_run(t); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    if size: r.font.size = Pt(size)
    r.font.name = font; return r


def add_para(doc, t="", bold=False, italic=False, color=None, size=10,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, sb=0, sa=5, keep=False):
    p = doc.add_paragraph(); p.alignment = align
    p.paragraph_format.space_before = Pt(sb); p.paragraph_format.space_after = Pt(sa)
    p.paragraph_format.keep_with_next = keep
    if t: add_run(p, t, bold, italic, color, size)
    return p


def add_heading(doc, t, level=1):
    return add_para(doc, t, bold=True, color=NAVY, size=(13 if level == 1 else 11),
                    align=WD_ALIGN_PARAGRAPH.LEFT, sb=(10 if level == 1 else 7), sa=3, keep=True)


def add_caption(doc, t):
    return add_para(doc, t, italic=True, color=RGBColor(0x40, 0x40, 0x40), size=9,
                    align=WD_ALIGN_PARAGRAPH.CENTER, sb=2, sa=8)


def add_pic(doc, path, w):
    doc.add_picture(path, width=Inches(w))
    pp = doc.paragraphs[-1]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.paragraph_format.keep_with_next = True   # keep figure with its caption


def make_table(doc, headers, rows_, caption, widths, hl=None):
    tbl = doc.add_table(rows=1 + len(rows_), cols=len(headers))
    tbl.style = "Table Grid"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, w in enumerate(widths):
        for row in tbl.rows:
            row.cells[i].width = Inches(w)
    for c in tbl.rows[0].cells: set_cell_bg(c, NAVY)
    for ci, h in enumerate(headers):
        pp = tbl.rows[0].cells[ci].paragraphs[0]; pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(pp, h, bold=True, color=WHITE, size=9)
    for ri, rd in enumerate(rows_):
        dr = tbl.rows[ri + 1]
        bg = HL if hl == ri else (WHITE if ri % 2 == 0 else GRAY_L)
        for c in dr.cells: set_cell_bg(c, bg)
        for ci, v in enumerate(rd):
            pp = dr.cells[ci].paragraphs[0]
            pp.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
            add_run(pp, str(v), size=9, bold=(hl == ri))
            pp.paragraph_format.keep_with_next = True   # keep the table and its caption together
    for c in tbl.rows[0].cells: c.paragraphs[0].paragraph_format.keep_with_next = True
    add_caption(doc, caption)


def add_apa(doc, t):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.5); p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.space_after = Pt(6); add_run(p, t, size=10)


def clear_body_after_table1(doc):
    body = doc.element.body; sect = body.find(qn("w:sectPr")); c = 0; cut = None
    for i, ch in enumerate(list(body)):
        if ch.tag == qn("w:tbl"):
            c += 1
            if c == 2: cut = i; break
    if cut is not None:
        for el in [e for e in list(body)[cut + 1:] if e.tag != qn("w:sectPr")]: body.remove(el)
    if sect is not None and sect not in list(body): body.append(sect)


# ============================ BUILD =========================================
def build():
    os.makedirs(OUT, exist_ok=True)
    p_imp, p_ser, p_mat = fig_imputacion(), fig_serie(), fig_matriz()
    p_3d, p_cf, p_hr = fig_planos(), fig_coef(), fig_horario()

    doc = Document(TEMPLATE)
    hdr = doc.tables[0].rows[0].cells[1]
    hdr.paragraphs[2].runs[0].text = "Maestría en Ciencias de los Datos | 3er Semestre"
    hdr.paragraphs[3].runs[0].text = "Materia: Visualización de Grandes Bases de Datos"
    tp = doc.paragraphs[1]; tp.clear()
    r = tp.add_run("Tarea 1.6. Imputación del NO₂ y exploración de su relación con el PM10 y el PM2.5\n"
                   "Estación de monitoreo Centro, Zona Metropolitana de Guadalajara, 2024")
    r.bold = True; r.font.size = Pt(13); r.font.color.rgb = NAVY; r.font.name = "Calibri"
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def si(idx, text):
        cell = doc.tables[1].rows[idx].cells[1]; cell.paragraphs[0].runs[0].text = text
        for extra in cell.paragraphs[0].runs[1:]: extra.text = ""
    si(1, "13 de octubre de 2026")
    si(2, "Tarea 1.6. Exploración entre NO₂ y PM2.5/PM10")
    si(3, "Python 3.13 (pandas, NumPy, Matplotlib, statsmodels)")
    si(4, "SIMAJ, base de calidad del aire 2024 (BD_2024)")
    si(5, "Equipo personal, entorno conda MCD")
    si(6, "Calidad del aire ZMG 2024, estación Centro (CEN)")
    clear_body_after_table1(doc)

    c14, c15, cc, ca = ev["1.4"], ev["1.5"], ev["comb"], ev["acot"]
    ene = mens.loc[1, "comp"]

    # ---------------- 1. Introduccion
    add_para(doc, sb=8, sa=0)
    add_heading(doc, "1. Introducción")
    add_para(doc,
        "En la estación Centro del SIMAJ el dióxido de nitrógeno (NO₂) solo se midió de septiembre "
        f"a diciembre de 2024. De las {num(n_total)} horas del año, a {num(n_falt)} les falta el dato "
        f"({pct(n_falt / n_total)}). En esta tarea primero completé esos valores con los modelos de "
        "regresión lineal simple de las tareas 1.4 y 1.5 y después, con la serie ya completa, revisé "
        "cómo se relaciona el NO₂ con las partículas PM10 y PM2.5. Para la segunda parte seguí lo "
        "visto en la sesión 8: gráficas del NO₂ contra cada partícula, gráficas de las partículas "
        "entre sí para buscar colinealidad y modelos de regresión múltiple, que se pueden ver como "
        "un plano o una superficie.")

    # ---------------- 2. Imputacion
    add_heading(doc, "2. Imputación de los valores faltantes de NO₂")
    add_heading(doc, "2.1 Elección de la variable explicativa", level=2)
    add_para(doc,
        "El NOx fue el contaminante con mayor correlación en las tareas anteriores, pero no sirve "
        "para imputar. Por un lado incluye al NO₂ en su cálculo, y por otro le faltan las mismas "
        "horas, ya que los dos salen del mismo analizador. Por eso usé el CO, que fue el siguiente "
        "con mayor correlación (r = 0.909 en la Tarea 1.5) y que se midió casi todo el año. Con él "
        f"se pueden completar {num(n_falt_co)} de las {num(n_falt)} horas faltantes. Las otras "
        f"{num(n_falt_sin)} tampoco tienen CO y las dejé vacías, porque estimarlas a partir de otra "
        "estimación ya no tendría mucho sentido.")
    add_para(doc,
        f"Revisé también las demás estaciones. De las {n_est} que tiene la base, {n_est_no2} miden "
        f"NO₂, y en ninguna hay datos antes de {MESES[mes_ini_red - 1]}. El hueco de enero a julio es "
        "de toda la red, así que no hay una estación vecina contra la cual comparar los valores "
        "imputados.")

    add_heading(doc, "2.2 Comparación de los modelos de las tareas 1.4 y 1.5", level=2)
    add_para(doc,
        "Comparé los dos modelos simples del NO₂ en función del CO. El de la Tarea 1.4 tiene "
        f"intercepto y se ajustó con todas las horas (NO₂ = {b0_14:.4f} + {b1_14:.4f}·CO). El de la "
        f"Tarea 1.5 pasa por el origen y se ajustó sin atípicos (NO₂ = {b_15:.4f}·CO). En la Tabla 1 "
        "están sus errores en las horas donde sí hay NO₂, separando las horas con CO por debajo y "
        f"por encima de la cerca de Tukey ({co_lim:.3f} ppm), que marca el rango con el que se "
        "estimó el modelo 1.5.", keep=True)
    erow = []
    for k, nm in [("1.4", "Modelo 1.4 (con intercepto)"), ("1.5", "Modelo 1.5 (por el origen)"),
                  ("comb", "Regla combinada (1.5 y 1.4)"), ("acot", "Modelo 1.5 acotado")]:
        e = ev[k]
        erow.append([nm, f(e["todo"]["rmse"], 4), f(e["bajo"]["rmse"], 4), f(e["alto"]["rmse"], 4),
                     f(e["alto"]["sesgo"], 4), f(oos[k]["rmse"], 4)])
    make_table(doc, ["Modelo", "RMSE total", f"RMSE CO ≤ {co_lim:.2f}", f"RMSE CO > {co_lim:.2f}",
                     f"Sesgo CO > {co_lim:.2f}", "RMSE nov-dic*"], erow,
               f"Tabla 1. Error de imputación (ppm) sobre las {num(len(p14))} horas con NO₂ y CO "
               f"observados; {num(n_alto_obs)} de ellas tienen CO por encima de la cerca. "
               "*Validación fuera de muestra: modelos reestimados con septiembre y octubre y "
               "evaluados en noviembre y diciembre. Sesgo = observado − estimado.",
               [2.1, 0.8, 1.0, 1.0, 1.0, 0.9], hl=3)
    add_para(doc,
        f"Con CO normal, el modelo 1.5 se equivoca menos (RMSE de {f(c15['bajo']['rmse'], 4)} contra "
        f"{f(c14['bajo']['rmse'], 4)} ppm). El problema aparece cuando el CO pasa de la cerca. Ahí "
        "la recta sigue subiendo, pero el NO₂ real ya no, porque la relación se aplana en valores "
        f"altos (lo mismo que se vio en la Tarea 1.4), y el modelo termina sobrestimando el NO₂ en "
        f"{f(-c15['alto']['sesgo'], 3)} ppm en promedio. El modelo 1.4 tiene menos pendiente y en "
        "ese tramo falla mucho menos.")
    add_para(doc,
        "Con eso en mente probé dos reglas más. La primera combina los modelos: el 1.5 hasta la "
        "cerca y el 1.4 de ahí en adelante. La segunda usa solo el 1.5, pero sin dejar que se salga "
        "de su rango: si el CO rebasa la cerca, el NO₂ se queda en el valor que el modelo da justo "
        f"en ella ({f(pred15(co_lim), 4)} ppm).", keep=True)
    add_pic(doc, p_imp, 5.6)
    add_caption(doc, "Figura 1. Modelos de regresión simple del NO₂ en función del CO y reglas de "
                     "imputación evaluadas. La línea azul gruesa es la regla adoptada.")
    add_para(doc,
        "Medida con los mismos datos del ajuste, la regla combinada sale un poco mejor (RMSE de "
        f"{f(cc['todo']['rmse'], 4)} contra {f(ca['todo']['rmse'], 4)} ppm). Sin embargo, tiene un "
        f"salto en la cerca: el NO₂ estimado baja de golpe de {f(pred15(co_lim), 3)} a "
        f"{f(pred14(co_lim), 3)} ppm, lo cual no tiene explicación física. Para decidir hice una "
        "prueba fuera de muestra: reajusté los modelos solo con septiembre y octubre y los apliqué "
        "a noviembre y diciembre. Las tres opciones que prolongan una recta quedaron con errores "
        f"parecidos, alrededor de {f(oos['comb']['rmse'], 3)} ppm, y el modelo acotado bajó a "
        f"{f(oos['acot']['rmse'], 4)} ppm. Por eso me quedé con el modelo 1.5 acotado. La prueba "
        "también deja ver que la relación entre el NO₂ y el CO no es igual todos los meses, así que "
        "los valores imputados hay que leerlos como una aproximación.")

    add_heading(doc, "2.3 Serie completada", level=2)
    add_para(doc,
        f"En total se imputaron {num(n_falt_co)} horas. En {num(n_imp_alto)} de ellas el CO estaba "
        "por encima de la cerca y el NO₂ se quedó en el valor límite. La Figura 2 muestra la serie "
        "diaria y la Tabla 2 los promedios por mes.", keep=True)
    add_pic(doc, p_ser, 6.4)
    add_caption(doc, "Figura 2. Promedio diario de NO₂ observado (septiembre a diciembre) e "
                     "imputado a partir del CO, estación Centro, 2024.")
    mrow = [[MESES[m - 1].capitalize(), num(mens.loc[m, "n_obs"]), num(mens.loc[m, "n_imp"]),
             "" if np.isnan(mens.loc[m, "obs"]) else f(mens.loc[m, "obs"], 4), f(mens.loc[m, "comp"], 4)]
            for m in range(1, 13)]
    make_table(doc, ["Mes", "Horas observadas", "Horas imputadas", "NO₂ observado", "NO₂ completado"],
               mrow, "Tabla 2. Horas observadas e imputadas y promedio mensual de NO₂ (ppm).",
               [1.3, 1.1, 1.1, 1.1, 1.1])
    add_para(doc,
        "De septiembre a diciembre los promedios casi no se mueven, porque ahí solo se rellenaron "
        f"{num(n_falt_sd)} horas sueltas. De enero a agosto, en cambio, todo viene del CO. El "
        f"resultado se ve lógico: enero tiene el promedio más alto ({f(ene, 4)} ppm) y los valores "
        "bajan hacia el verano. Eso cuadra con lo que pasa en Guadalajara, donde en los meses secos "
        "las inversiones térmicas atrapan los contaminantes y en la temporada de lluvias se "
        "dispersan mejor. Aun así, este patrón viene del CO y no hay mediciones de NO₂ en esos "
        "meses para confirmarlo.")

    # ---------------- 3. Exploracion
    add_heading(doc, "3. Relación del NO₂ con el PM10 y el PM2.5")
    add_heading(doc, "3.1 Las partículas PM10 y PM2.5", level=2)
    add_para(doc,
        "El material particulado (PM) es la mezcla de partículas sólidas y gotitas que quedan "
        "suspendidas en el aire. Se clasifica por tamaño. El PM10 son las partículas de 10 µm o "
        "menos y el PM2.5 las de 2.5 µm o menos, unas treinta veces más delgadas que un cabello. "
        "El PM2.5 forma parte del PM10, y lo que queda entre los dos (partículas de 2.5 a 10 µm) se "
        "conoce como fracción gruesa. A diferencia del NO₂ y del CO, que son gases, las partículas "
        "se miden en µg/m³ (Seinfeld y Pandis, 2016).")
    add_para(doc,
        "Las dos fracciones vienen de fuentes distintas. La gruesa se genera sobre todo por "
        "procesos mecánicos, como el polvo que levantan el viento y los coches, la construcción o "
        "el desgaste de llantas y frenos. El PM2.5 sale principalmente de la combustión, en "
        "especial de motores diésel, aunque una parte se forma en la propia atmósfera. Ahí entra "
        "el NO₂, que se oxida a ácido nítrico y, al reaccionar con amoniaco, forma nitrato de "
        "amonio, que termina como PM2.5 (Seinfeld y Pandis, 2016). Por lo tanto, se esperaría "
        "encontrar alguna relación entre el NO₂ y las partículas finas, tanto por la fuente "
        "vehicular que comparten como por esta química.")
    add_para(doc,
        "En la salud, el tamaño importa. El PM10 llega hasta los bronquios y el PM2.5 entra a los "
        "alvéolos y puede pasar a la sangre, por lo que se asocia con enfermedades respiratorias y "
        "cardiovasculares. La OMS (2021) recomienda no pasar de 45 µg/m³ de PM10 ni de 15 µg/m³ de "
        "PM2.5 en promedio de 24 horas. En la estación Centro, en 2024, el PM10 promedió "
        f"{f(pm['PM10']['media'], 1)} µg/m³ y el PM2.5 {f(pm['PM2.5']['media'], 1)} µg/m³. El "
        f"promedio diario rebasó la recomendación en {pm['PM10']['exc']} de {pm['PM10']['dias']} días "
        f"para el PM10 y en {pm['PM2.5']['exc']} de {pm['PM2.5']['dias']} para el PM2.5. Además, la "
        f"mediana del cociente PM2.5/PM10 fue de {f(ratio_pm, 2)}; es decir, en esta estación algo "
        "menos de la mitad de la masa de partículas es fracción fina.")

    add_heading(doc, "3.2 Corpus de análisis", level=2)
    add_para(doc,
        "Tomé las horas que tienen NO₂ (medido o imputado), PM10 y PM2.5 al mismo tiempo, que son "
        f"{num(len(Zc_all))}. Igual que en la Tarea 1.5, quité las horas atípicas en cualquiera de "
        f"las tres variables con el criterio de Tukey, y quedaron {num(len(Zc))}, de las cuales "
        f"{num((Zc['origen'] == 'imputado').sum())} tienen NO₂ imputado. Para tener una referencia, "
        "repetí todos los modelos usando solo las horas con NO₂ medido "
        f"({num(len(Zo))} después de la misma limpieza).")

    add_heading(doc, "3.3 Dispersión entre las variables y colinealidad", level=2)
    add_para(doc,
        "Siguiendo la sesión 8, la Figura 3 muestra el NO₂ contra cada partícula y, en el tercer "
        "panel, las dos partículas entre sí. Los coeficientes de correlación están en la Tabla 3.",
        keep=True)
    add_pic(doc, p_mat, 6.5)
    add_caption(doc, "Figura 3. Dispersión del NO₂ frente al PM10 y al PM2.5, y del PM10 frente al "
                     "PM2.5, serie completada sin atípicos. Los puntos ámbar son horas con NO₂ imputado.")
    rc_, ro_ = Mc["r"], Mo["r"]
    make_table(doc, ["Par de variables", "r serie completada", "r solo observadas"],
               [["NO₂ – PM10", f(rc_.loc["NO2", "PM10"]), f(ro_.loc["NO2", "PM10"])],
                ["NO₂ – PM2.5", f(rc_.loc["NO2", "PM2.5"]), f(ro_.loc["NO2", "PM2.5"])],
                ["PM10 – PM2.5", f(rc_.loc["PM10", "PM2.5"]), f(ro_.loc["PM10", "PM2.5"])],
                ["VIF de PM10 y PM2.5", f(Mc["vif"], 2), f(Mo["vif"], 2)]],
               "Tabla 3. Correlación de Pearson entre el NO₂ y las partículas, y entre las partículas, "
               "corpus sin atípicos.", [2.2, 1.4, 1.4], hl=2)
    add_para(doc,
        "El NO₂ tiene una correlación positiva y moderada con las dos partículas, un poco mayor con "
        "el PM10. Lo que más destaca, sin embargo, es la correlación entre las partículas mismas "
        f"(r = {f(rc_.loc['PM10', 'PM2.5'])} en la serie completada y {f(ro_.loc['PM10', 'PM2.5'])} "
        "con los datos medidos). Se explica porque el PM10 contiene al PM2.5, y significa que las "
        "dos variables independientes traen en buena parte la misma información, que es la "
        "colinealidad que hay que revisar antes de un modelo múltiple. El VIF, calculado como "
        f"1 / (1 − R²ⱼ), fue de {f(Mc['vif'], 2)} y {f(Mo['vif'], 2)}. No llega a 5 o 10, que es "
        "donde normalmente se considera grave, pero sí quiere decir que la varianza de cada "
        "coeficiente se multiplica por tres o cuatro.")

    add_heading(doc, "3.4 Modelos de regresión lineal múltiple", level=2)
    add_para(doc,
        "Ajusté cuatro modelos por mínimos cuadrados: los dos simples, el múltiple "
        "NO₂ = b₀ + b₁·PM10 + b₂·PM2.5 y uno cuadrático en el PM10, "
        "NO₂ = b₀ + b₁·PM10 + b₂·PM10² + b₃·PM2.5, parecido al ejemplo de la sesión 8. Los "
        "resultados están en la Tabla 4 y los planos en la Figura 4.", keep=True)

    def fila(M, k, nm):
        m = M[k]; pr, pv = m.params, m.pvalues
        b10 = f"{pr['PM10'] * 1000:.3f}" if "PM10" in pr else "–"
        b25 = f"{pr['PM2.5'] * 1000:.3f}" if "PM2.5" in pr else "–"
        if "PM2.5" in pr and pv["PM2.5"] > 0.05:
            b25 += f" (p = {pv['PM2.5']:.2f})"
        bq = f"{pr['PM10_2'] * 1000:.5f} (p = {pv['PM10_2']:.2f})" if "PM10_2" in pr else "–"
        return [nm, b10, bq, b25, f(m.rsquared), f(m.rsquared_adj)]

    trow = []
    for M, tag in [(Mc, "Completada"), (Mo, "Observada")]:
        for k, nm in [("PM10", "simple PM10"), ("PM2.5", "simple PM2.5"), ("lineal", "múltiple lineal"),
                      ("cuad", "múltiple cuadrático")]:
            trow.append(fila(M, k, f"{tag}: {nm}"))
    make_table(doc, ["Serie y modelo", "b PM10", "b PM10²", "b PM2.5", "R²", "R² ajustado"], trow,
               "Tabla 4. Coeficientes (ppb de NO₂ por µg/m³; por (µg/m³)² en el término cuadrático) "
               "y ajuste de los modelos. Salvo donde se indica el valor p, los coeficientes son "
               "significativos al 5 %.", [2.2, 0.7, 1.4, 1.2, 0.5, 0.8], hl=2)
    add_pic(doc, p_3d, 6.6)
    add_caption(doc, "Figura 4. Plano de regresión lineal múltiple y superficie del modelo cuadrático "
                     "del NO₂ en función del PM10 y el PM2.5, serie completada (muestra de 1,500 horas).")

    lc, lo_ = Mc["lineal"], Mo["lineal"]
    b10o = Mo["PM10"].params["PM10"]; b25o = Mo["PM2.5"].params["PM2.5"]
    add_para(doc,
        "Cada coeficiente indica cuánto cambia el NO₂ por cada µg/m³ más de partículas. Con los "
        f"datos medidos, 10 µg/m³ más de PM10 van con {f(b10o * 10000, 1)} ppb más de NO₂, y 10 µg/m³ "
        f"más de PM2.5 con {f(b25o * 10000, 1)} ppb. Como referencia, el NO₂ promedio en esas horas "
        f"es de {f(no2_med_o * 1000, 1)} ppb. El coeficiente del PM2.5 sale más grande porque sus "
        "niveles son más bajos y 10 µg/m³ pesan más en términos relativos. Esto no quiere decir que "
        "las partículas generen NO₂. Lo que pasa es que los dos suben juntos, porque tienen fuentes "
        "de combustión en común y se acumulan con el mismo clima, como las inversiones térmicas o "
        "los días sin viento.")
    add_para(doc,
        "En los datos medidos, meter el PM2.5 a un modelo que ya tiene al PM10 no mejora nada "
        f"(R² de {f(Mo['PM10'].rsquared)} en los dos casos) y su coeficiente deja de ser "
        f"significativo (p = {lo_.pvalues['PM2.5']:.2f}), aunque en el modelo simple sí lo era. Es "
        "el efecto clásico de la colinealidad. En la Figura 5 se ve bien: el coeficiente del PM2.5 "
        f"baja de {f(Mo['PM2.5'].params['PM2.5'] * 1000)} a {f(lo_.params['PM2.5'] * 1000)} ppb por "
        "µg/m³ y su intervalo de confianza toca el cero.", keep=True)
    add_pic(doc, p_cf, 4.8)
    add_caption(doc, "Figura 5. Coeficiente del PM2.5 en el modelo simple y en el múltiple, con su "
                     "intervalo de confianza de 95 %, para la serie observada y la completada.")
    add_para(doc,
        "En la serie completada pasa algo distinto: los dos coeficientes son significativos y el "
        f"R² sube de {f(Mc['PM10'].rsquared)} a {f(lc.rsquared)}. Creo que esto se debe sobre todo "
        "a la imputación. En las horas imputadas el NO₂ no es una medición sino una función del CO, "
        "así que en esos meses el modelo está midiendo en realidad cómo se relaciona el CO con las "
        "partículas. Como el CO también viene de la combustión, se relaciona de forma más pareja "
        "con el PM10 y el PM2.5, y eso le da al PM2.5 un peso que no tiene en los datos medidos.")
    add_para(doc,
        "El término cuadrático no aportó nada en ninguna serie "
        f"(p = {Mc['cuad'].pvalues['PM10_2']:.2f} y {Mo['cuad'].pvalues['PM10_2']:.2f}), y en la "
        "Figura 4 la superficie casi no se curva. Con el NOx y el CO sí había curvatura, pero con las "
        "partículas, ya sin atípicos, la relación es prácticamente lineal. Lo que limita el modelo "
        f"es la dispersión de los puntos: entre las dos partículas explican {pct(lc.rsquared)} de la "
        "variabilidad del NO₂ en la serie completada.")

    add_heading(doc, "3.5 Fracción fina y fracción gruesa", level=2)
    mo, mc = Go["m"], Gc["m"]
    add_para(doc,
        "Como la colinealidad viene de que el PM10 contiene al PM2.5, se puede quitar separando el "
        "PM10 en sus dos partes, la fracción fina (PM2.5) y la gruesa (PM10 − PM2.5), y ajustando "
        "NO₂ = b₀ + b₁·PM2.5 + b₂·(PM10 − PM2.5). El ajuste es el mismo que el del modelo múltiple de "
        "la Tabla 4, porque la información es la misma, pero ahora las dos variables ya no "
        f"comparten partículas. Su correlación baja a {f(Go['r'])} en los datos medidos y el VIF a "
        f"{f(Go['vif'], 2)} (Tabla 5).", keep=True)
    ci_o, ci_c = mo.conf_int(), mc.conf_int()

    def ic(ci, v):
        return f"{ci.loc[v, 0] * 1000:.3f} a {ci.loc[v, 1] * 1000:.3f}"
    make_table(doc, ["Serie", "b fracción fina (IC 95 %)", "b fracción gruesa (IC 95 %)", "R²", "VIF"],
               [["Observada", f"{mo.params['PM2.5'] * 1000:.3f} ({ic(ci_o, 'PM2.5')})",
                 f"{mo.params['PMg'] * 1000:.3f} ({ic(ci_o, 'PMg')})", f(mo.rsquared), f(Go["vif"], 2)],
                ["Completada", f"{mc.params['PM2.5'] * 1000:.3f} ({ic(ci_c, 'PM2.5')})",
                 f"{mc.params['PMg'] * 1000:.3f} ({ic(ci_c, 'PMg')})", f(mc.rsquared), f(Gc["vif"], 2)]],
               "Tabla 5. Modelo del NO₂ con la fracción fina y la fracción gruesa de las partículas "
               "(ppb de NO₂ por µg/m³). Todos los coeficientes son significativos al 1 %.",
               [1.1, 2.0, 2.0, 0.6, 0.6], hl=0)
    add_para(doc,
        "Así, los dos coeficientes salen significativos y con intervalos estrechos. En los datos "
        f"medidos, cada µg/m³ de partículas finas se asocia con {f(mo.params['PM2.5'] * 1000)} ppb de "
        f"NO₂ y cada µg/m³ de gruesas con {f(mo.params['PMg'] * 1000)} ppb. Las dos fracciones se "
        "relacionan con el NO₂ más o menos igual, con algo de ventaja para la fina, lo cual tiene "
        "sentido si se piensa que viene de la combustión.")
    add_para(doc,
        "Esto cambia la forma de leer la sección anterior. Que el PM2.5 dejara de ser significativo "
        "junto al PM10 no quería decir que las partículas finas no tengan que ver con el NO₂, sino que "
        "no tienen mucho más que ver que las gruesas. En el modelo con PM10, el coeficiente del "
        "PM2.5 mide precisamente esa diferencia. En la serie completada la fracción fina sale con "
        "más ventaja, otra vez por efecto de la imputación con CO.")
    if Go["neg"] or Gc["neg"]:
        add_para(doc,
            f"Un detalle: en {Go['neg']} horas medidas y {Gc['neg']} de la serie completada el PM2.5 "
            "aparece un poco mayor que el PM10, cosa que físicamente no puede pasar y que atribuyo a "
            "la incertidumbre de los equipos. Como son muy pocas, las dejé.")

    add_heading(doc, "3.6 Ciclo diario", level=2)
    pk = hor["NO2"].idxmax(); mn = hor["NO2"].idxmin()
    add_para(doc,
        "Por último, quise ver si los contaminantes siguen el mismo horario a lo largo del día. En "
        "la Figura 6 está el promedio de cada hora dividido entre el promedio general de cada "
        "contaminante, para poder compararlos en la misma escala. Usé solo las horas con NO₂ medido, "
        "para que la imputación no influyera.", keep=True)
    add_pic(doc, p_hr, 6.2)
    add_caption(doc, "Figura 6. Ciclo diario promedio del NO₂, el PM10, el PM2.5 y el CO, expresado "
                     "como proporción del promedio de cada variable, estación Centro, septiembre a "
                     "diciembre de 2024.")
    add_para(doc,
        "Los cuatro siguen casi el mismo patrón. Suben en la mañana, hacia las 7 u 8, bajan durante "
        "el mediodía y la tarde, y vuelven a subir desde las 18 o 19 horas hasta la noche. El NO₂ "
        f"llega a su mínimo a las {mn} horas y a su máximo a las {pk}, cuando es "
        f"{f(hor['NO2'].max() / hor['NO2'].min(), 1)} veces mayor. Los picos coinciden con las horas "
        "de más tráfico. La baja del mediodía se debe a que el sol calienta el suelo y la capa de "
        "mezcla de la atmósfera crece, lo que diluye los contaminantes. En el NO₂ la caída es más "
        "fuerte que en las partículas porque a esas horas, además, se consume en las reacciones que "
        "forman ozono.")
    add_para(doc,
        "Este patrón compartido ayuda a entender la correlación entre el NO₂ y las partículas. Los "
        "dos responden al mismo tráfico y a las mismas condiciones de dispersión, sin que uno sea "
        "causa del otro.")

    # ---------------- 4. Conclusiones
    add_heading(doc, "4. Conclusiones")
    add_para(doc,
        f"Con el CO pude completar {pct(n_falt_co / n_falt)} de las horas sin NO₂. Ninguno de los "
        "dos modelos anteriores funcionaba bien en todo el rango. La recta por el origen de la "
        "Tarea 1.5 es la más precisa con valores normales de CO, pero cuando el CO es alto "
        "sobrestima el NO₂, porque se sale del rango con que se ajustó. Limitarla a ese rango fue la "
        "opción que mejor aguantó la prueba fuera de muestra, con menos de la mitad del error de "
        "las otras. Aun así, la relación entre los dos contaminantes cambia de mes a mes, y los "
        "ocho meses imputados no se pueden comparar con mediciones reales.")
    add_para(doc,
        "El NO₂ tiene una relación positiva y moderada con las partículas, algo mayor con el PM10. "
        "La correlación más fuerte de todo el análisis, sin embargo, fue entre el PM10 y el PM2.5, "
        "y esa colinealidad hizo que, con los datos medidos, el PM2.5 dejara de ser significativo "
        "al agregar el PM10. Al separar el PM10 en fracción fina y gruesa el problema desaparece, y "
        "las dos fracciones resultan significativas y con un efecto parecido sobre el NO₂, un poco "
        "mayor en la fina.")
    add_para(doc,
        "El ciclo diario explica buena parte de lo anterior, ya que el NO₂ y las partículas suben "
        "con el tráfico de la mañana y de la noche y bajan juntos al mediodía. Aun así, las "
        "partículas explican mucho menos del NO₂ que el CO o el NOx. Me parece razonable, porque "
        "además de la combustión tienen otras fuentes, como el polvo y la construcción, mientras "
        "que el NO₂ depende casi solo del tráfico.")

    add_heading(doc, "Referencias")
    add_apa(doc, "Little, R. J. A., & Rubin, D. B. (2019). Statistical analysis with missing data "
                 "(3.ª ed.). Wiley.")
    add_apa(doc, "Montgomery, D. C., Peck, E. A., & Vining, G. G. (2021). Introduction to linear "
                 "regression analysis (6.ª ed.). Wiley.")
    add_apa(doc, "Organización Mundial de la Salud. (2021). WHO global air quality guidelines: "
                 "Particulate matter (PM2.5 and PM10), ozone, nitrogen dioxide, sulfur dioxide and "
                 "carbon monoxide. World Health Organization. https://iris.who.int/handle/10665/345329")
    add_apa(doc, "Ruiz Chávez, C. F. (2026). Visualización de grandes bases de datos: Sesión 8 "
                 "[Diapositivas de clase]. Maestría en Ciencia de los Datos, Centro Universitario de "
                 "Ciencias Económico Administrativas, Universidad de Guadalajara.")
    add_apa(doc, "Secretaría de Medio Ambiente y Desarrollo Territorial. (2024). Sistema de "
                 "Monitoreo Atmosférico de Jalisco (SIMAJ): base de datos de calidad del aire 2024. "
                 "Gobierno del Estado de Jalisco. https://siga.jalisco.gob.mx/aire/")
    add_apa(doc, "Seinfeld, J. H., & Pandis, S. N. (2016). Atmospheric chemistry and physics: From "
                 "air pollution to climate change (3.ª ed.). Wiley.")
    out = OUT + r"\Tarea1_6_VisualizacionGBD_CarlosPulidoRosas.docx"
    doc.save(out); print("DOCX:", out)
    subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", OUT, out],
                   check=True, capture_output=True)
    print("PDF:", out.replace(".docx", ".pdf"))
    return out


if __name__ == "__main__":
    build()
