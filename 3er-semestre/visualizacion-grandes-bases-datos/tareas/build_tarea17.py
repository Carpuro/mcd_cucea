# -*- coding: utf-8 -*-
"""
Tarea 1.7 - Refined multiple regression NO2 ~ PM10 + PM2.5: line forced through the origin and
Tukey outliers of NO2, PM10 and PM2.5 removed, compared against the Tarea 1.6 models (intercept).
Also looks at how PM10 and PM2.5 relate to each other, with K-means as in Session 9.
Reuses the data preparation (imputed NO2 series, Tukey-clean corpus) of build_tarea16.
"""
import os
import sys
import subprocess
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import statsmodels.api as sm
from sklearn.cluster import KMeans
from sklearn.preprocessing import StandardScaler
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

os.environ.setdefault("LOKY_MAX_CPU_COUNT", "4")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_tarea16 as t16                                     # noqa: E402
from build_tarea16 import (add_para, add_heading, add_caption, add_pic, make_table,  # noqa: E402
                           add_apa, clear_body_after_table1, NAVY, NAVY_HEX, AMBER_HEX,
                           RED_HEX, TEMPLATE, SOFFICE, f, pct, num)

COURSE = r"C:\Users\carlo\GitHub\mcd_cucea\3er-semestre\visualizacion-grandes-bases-datos"
OUT = COURSE + r"\tareas\Tarea1_7"
GRAY_HEX = "#777777"


# ============================ MODELS ========================================
def ajuste(Z, origen):
    X = Z[["PM10", "PM2.5"]]
    X = X if origen else sm.add_constant(X)
    m = sm.OLS(Z["NO2"], X).fit()
    y, e = Z["NO2"], m.resid
    return {"m": m, "n": len(Z), "r2_modelo": m.rsquared,
            "r2_media": 1 - (e ** 2).sum() / ((y - y.mean()) ** 2).sum(),
            "rmse": np.sqrt((e ** 2).mean()), "sesgo": e.mean(),
            "r_y_yhat": np.corrcoef(y, m.fittedvalues)[0, 1]}


SERIES = {"Completada": (t16.Zc_all, t16.Zc), "Observada": (t16.Zo_all, t16.Zo)}
M = {}
for s, (ZA, Z) in SERIES.items():
    M[s] = {"bruto_int": ajuste(ZA, False), "t16": ajuste(Z, False),
            "t17": ajuste(Z, True), "bruto_origen": ajuste(ZA, True),
            "r_bruto": ZA[t16.VARS].corr(), "r_limpio": Z[t16.VARS].corr(),
            "n_bruto": len(ZA), "n_limpio": len(Z)}
    # PM2.5 as a fraction of PM10 (line through the origin)
    mp = sm.OLS(Z["PM2.5"], Z[["PM10"]]).fit()
    M[s]["frac"] = mp.params.iloc[0]; M[s]["frac_r2"] = mp.rsquared
    M[s]["vif"] = 1 / (1 - Z[["PM10", "PM2.5"]].corr().iloc[0, 1] ** 2)

C, O = M["Completada"], M["Observada"]

# K-means on the observed hours (standardized PM10 and PM2.5)
Zk = t16.Zo.copy()
Xk = StandardScaler().fit_transform(Zk[["PM10", "PM2.5"]].values)
inercias = {k: KMeans(k, n_init=10, random_state=42).fit(Xk).inertia_ for k in range(1, 8)}
K = 3
km = KMeans(K, n_init=10, random_state=42).fit(Xk)
Zk["grupo"] = km.labels_
# relabel the groups from low to high PM10 so the names make sense
orden = Zk.groupby("grupo")["PM10"].mean().sort_values().index
Zk["grupo"] = Zk["grupo"].map({g: i for i, g in enumerate(orden)})
grupos = Zk.groupby("grupo").agg(n=("NO2", "size"), PM10=("PM10", "mean"), PM25=("PM2.5", "mean"),
                                  NO2=("NO2", "mean"))
grupos["ratio"] = grupos["PM25"] / grupos["PM10"]
NOMBRES_G = ["Bajo", "Medio", "Alto"]


# ============================ FIGURES =======================================
def fig_pred_vs_real():
    p = OUT + r"\real_vs_predicho.png"
    Z = SERIES["Completada"][1]
    fig, ax = plt.subplots(1, 2, figsize=(10.5, 4.4), sharey=True)
    for axx, key, titulo, col in [(ax[0], "t16", "Tarea 1.6: con intercepto", NAVY_HEX),
                                  (ax[1], "t17", "Tarea 1.7: por el origen", AMBER_HEX)]:
        r = C[key]
        axx.scatter(Z["NO2"], r["m"].fittedvalues, s=4, alpha=0.2, color=col, edgecolors="none")
        lim = [0, Z["NO2"].max() * 1.05]
        axx.plot(lim, lim, ls="--", color=GRAY_HEX, lw=1.2, label="predicción perfecta")
        axx.set_xlim(lim); axx.set_ylim(lim)
        axx.set_title(f"{titulo}\nRMSE = {r['rmse'] * 1000:.1f} ppb, sesgo = {r['sesgo'] * 1000:+.1f} ppb",
                      color=NAVY_HEX, weight="bold", fontsize=10)
        axx.set_xlabel("NO₂ real (ppm)", fontsize=9); axx.grid(linestyle=":", alpha=0.4)
        axx.legend(fontsize=8, loc="upper left")
    ax[0].set_ylabel("NO₂ estimado (ppm)", fontsize=9)
    plt.tight_layout(); plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


def fig_residuales():
    p = OUT + r"\residuales.png"
    fig, ax = plt.subplots(2, 2, figsize=(10.5, 7))
    for j, (key, titulo, col) in enumerate([("t16", "Tarea 1.6: con intercepto", NAVY_HEX),
                                            ("t17", "Tarea 1.7: por el origen", AMBER_HEX)]):
        r = C[key]; e = r["m"].resid; yh = r["m"].fittedvalues
        a = ax[0, j]
        a.scatter(yh, e, s=4, alpha=0.2, color=col, edgecolors="none")
        a.axhline(0, color="black", lw=1)
        # running mean of the residuals to make the pattern visible
        orden_ = np.argsort(yh.values)
        movil = pd.Series(e.values[orden_]).rolling(300, center=True).mean()
        a.plot(yh.values[orden_], movil, color=RED_HEX, lw=2, label="promedio móvil de residuales")
        a.set_title(f"Residuales vs. estimado: {titulo}", color=NAVY_HEX, weight="bold", fontsize=9.5)
        a.set_xlabel("NO₂ estimado (ppm)", fontsize=8); a.set_ylabel("Residual (ppm)", fontsize=8)
        a.legend(fontsize=7, loc="upper left"); a.grid(linestyle=":", alpha=0.4)
        b = ax[1, j]
        b.hist(e, bins=60, color=col, alpha=0.8)
        b.axvline(0, color="black", lw=1)
        b.axvline(e.mean(), color=RED_HEX, lw=2, ls="--", label=f"media = {e.mean() * 1000:+.1f} ppb")
        b.set_title(f"Distribución de residuales: {titulo}", color=NAVY_HEX, weight="bold", fontsize=9.5)
        b.set_xlabel("Residual (ppm)", fontsize=8); b.set_ylabel("Horas", fontsize=8)
        b.legend(fontsize=7); b.grid(linestyle=":", alpha=0.4)
    plt.tight_layout(); plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


def fig_pm():
    p = OUT + r"\relacion_pm10_pm25.png"
    fig, ax = plt.subplots(1, 2, figsize=(10.5, 4.3), gridspec_kw={"width_ratios": [1.6, 1]})
    cols = ["#7FB3D5", AMBER_HEX, RED_HEX]
    for g in range(K):
        s = Zk[Zk["grupo"] == g]
        ax[0].scatter(s["PM10"], s["PM2.5"], s=6, alpha=0.4, color=cols[g], edgecolors="none",
                      label=f"Grupo {NOMBRES_G[g].lower()} ({len(s)} h)")
    for g in range(K):
        ax[0].scatter(grupos.loc[g, "PM10"], grupos.loc[g, "PM25"], marker="+", s=220, color="black", lw=2)
    xs = np.linspace(0, Zk["PM10"].max(), 50)
    ax[0].plot(xs, O["frac"] * xs, color="black", lw=1.5, ls="--",
               label=f"PM2.5 = {O['frac']:.2f}·PM10")
    ax[0].set_xlim(left=0); ax[0].set_ylim(bottom=0)
    ax[0].set_xlabel("PM10 (µg/m³)", fontsize=9); ax[0].set_ylabel("PM2.5 (µg/m³)", fontsize=9)
    ax[0].set_title("PM2.5 vs. PM10 y grupos de K-medias (k = 3)", color=NAVY_HEX, weight="bold", fontsize=10)
    ax[0].legend(fontsize=7.5, loc="upper left"); ax[0].grid(linestyle=":", alpha=0.4)
    ks = list(inercias.keys())
    ax[1].plot(ks, [inercias[k] for k in ks], marker="o", color=NAVY_HEX)
    ax[1].axvline(K, color=AMBER_HEX, ls="--", lw=1.5)
    ax[1].set_xlabel("Número de grupos k", fontsize=9); ax[1].set_ylabel("J (suma de distancias²)", fontsize=9)
    ax[1].set_title("Criterio del codo", color=NAVY_HEX, weight="bold", fontsize=10)
    ax[1].grid(linestyle=":", alpha=0.4)
    plt.tight_layout(); plt.savefig(p, dpi=190, bbox_inches="tight"); plt.close()
    return p


# ============================ BUILD =========================================
def build():
    os.makedirs(OUT, exist_ok=True)
    p_rv, p_res, p_pm = fig_pred_vs_real(), fig_residuales(), fig_pm()

    doc = Document(TEMPLATE)
    hdr = doc.tables[0].rows[0].cells[1]
    hdr.paragraphs[2].runs[0].text = "Maestría en Ciencias de los Datos | 3er Semestre"
    hdr.paragraphs[3].runs[0].text = "Materia: Visualización de Grandes Bases de Datos"
    tp = doc.paragraphs[1]; tp.clear()
    r = tp.add_run("Tarea 1.7. Modelo refinado del NO₂ con PM10 y PM2.5\n"
                   "Estación de monitoreo Centro, Zona Metropolitana de Guadalajara, 2024")
    r.bold = True; r.font.size = Pt(13); r.font.color.rgb = NAVY; r.font.name = "Calibri"
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def si(idx, text):
        cell = doc.tables[1].rows[idx].cells[1]; cell.paragraphs[0].runs[0].text = text
        for extra in cell.paragraphs[0].runs[1:]: extra.text = ""
    si(1, "21 de octubre de 2026")
    si(2, "Tarea 1.7. Modelado de NO₂ y PM2.5/PM10 refinado")
    si(3, "Python 3.13 (pandas, NumPy, Matplotlib, statsmodels, scikit-learn)")
    si(4, "SIMAJ, base de calidad del aire 2024 (BD_2024)")
    si(5, "Equipo personal, entorno conda MCD")
    si(6, "Calidad del aire ZMG 2024, estación Centro (CEN)")
    clear_body_after_table1(doc)

    c16, c17, o16, o17 = C["t16"], C["t17"], O["t16"], O["t17"]

    # ---------------- 1. Introduccion
    add_para(doc, sb=8, sa=0)
    add_heading(doc, "1. Introducción")
    add_para(doc,
        "En la Tarea 1.6 completé la serie de NO₂ de la estación Centro con el CO y ajusté modelos de "
        "regresión múltiple del NO₂ con el PM10 y el PM2.5. Ahora toca refinar ese modelo: forzar el "
        "plano de regresión a pasar por el origen, NO₂ = b₁·PM10 + b₂·PM2.5, y quitar los valores "
        "atípicos de las tres variables. Después lo comparo contra los modelos de la 1.6 para "
        "responder cuatro preguntas: si mejoró la correlación, si mejoraron el R² y los residuales, "
        "cómo se relacionan las partículas entre sí y qué nos dice el modelo.")
    add_para(doc,
        "Uso las mismas dos series de la tarea anterior: la serie completada, que incluye el NO₂ "
        f"imputado ({num(C['n_bruto'])} horas con las tres variables), y la serie observada, solo con "
        f"NO₂ medido ({num(O['n_bruto'])} horas). La segunda sirve como control, porque no depende de "
        "la imputación.")

    # ---------------- 2. Metodologia
    add_heading(doc, "2. Metodología")
    add_heading(doc, "2.1 Exclusión de atípicos", level=2)
    add_para(doc,
        "Como en las tareas 1.5 y 1.6, usé el criterio de Tukey: una hora es atípica si el NO₂, el "
        "PM10 o el PM2.5 queda fuera de [Q1 − 1.5·RIC, Q3 + 1.5·RIC]. Si una hora es atípica en "
        "cualquiera de las tres, la quito completa. Así quedan "
        f"{num(C['n_limpio'])} horas en la serie completada y {num(O['n_limpio'])} en la observada. "
        "Para ver el efecto de este paso, también calculé todo con los datos sin depurar.")
    add_heading(doc, "2.2 Modelo por el origen", level=2)
    add_para(doc,
        "El modelo de la 1.6 tenía intercepto, NO₂ = b₀ + b₁·PM10 + b₂·PM2.5. El refinado lo quita, "
        "así que cuando las dos partículas valen cero el NO₂ estimado también es cero. Para comparar "
        "los dos modelos de forma justa no basta con el R² que reporta cada uno, porque en un modelo "
        "sin intercepto el R² se calcula contra cero y no contra el promedio, y por eso sale mucho "
        "más alto aunque el ajuste sea peor (lo mismo que vimos en la Tarea 1.5). Por eso comparo "
        "también el R² calculado contra el promedio en los dos modelos, el error cuadrático medio "
        "(RMSE) y el promedio de los residuales, que en un modelo sin sesgo debe ser cero.")
    add_heading(doc, "2.3 Relación entre las partículas y K-medias", level=2)
    add_para(doc,
        "Para la relación entre el PM10 y el PM2.5 uso su correlación, una recta por el origen "
        "PM2.5 = b·PM10 y, siguiendo la sesión 9, el algoritmo de K-medias. K-medias agrupa las horas "
        "en k grupos minimizando J = Σ Σ ‖xᵢ − cⱼ‖², la suma de las distancias euclidianas al "
        "cuadrado de cada punto a su centroide. Antes de agrupar estandaricé las dos variables para "
        "que pesaran igual, y elegí k con el criterio del codo.")

    # ---------------- 3. Resultados
    add_heading(doc, "3. Resultados")
    add_heading(doc, "3.1 ¿Mejoró la correlación?", level=2)
    add_para(doc,
        "La Tabla 1 compara las correlaciones antes y después de quitar los atípicos.", keep=True)
    rb_c, rl_c, rb_o, rl_o = C["r_bruto"], C["r_limpio"], O["r_bruto"], O["r_limpio"]
    pares = [("NO₂ – PM10", "NO2", "PM10"), ("NO₂ – PM2.5", "NO2", "PM2.5"), ("PM10 – PM2.5", "PM10", "PM2.5")]
    make_table(doc, ["Par", "Completada sin depurar", "Completada depurada", "Observada sin depurar",
                     "Observada depurada"],
               [[nm, f(rb_c.loc[a, b]), f(rl_c.loc[a, b]), f(rb_o.loc[a, b]), f(rl_o.loc[a, b])]
                for nm, a, b in pares],
               "Tabla 1. Correlación de Pearson antes y después de excluir los atípicos de NO₂, PM10 y PM2.5.",
               [1.4, 1.3, 1.3, 1.3, 1.3])
    add_para(doc,
        "No, la correlación no mejoró, bajó. En la serie observada la del NO₂ con el PM10 pasa de "
        f"{f(rb_o.loc['NO2', 'PM10'])} a {f(rl_o.loc['NO2', 'PM10'])}, y con el PM2.5 de "
        f"{f(rb_o.loc['NO2', 'PM2.5'])} a {f(rl_o.loc['NO2', 'PM2.5'])}. Es lo mismo que pasó con las "
        "partículas en la Tarea 1.5: los atípicos son horas de contaminación muy alta en las que "
        "suben todos los contaminantes al mismo tiempo, y esos pocos puntos lejanos inflaban la "
        "correlación. Sin ellos la relación se ve más débil, pero es más representativa de lo que "
        "pasa en un día normal. Forzar el modelo por el origen no cambia esto, porque la correlación "
        "de Pearson no depende del modelo que se ajuste.")

    add_heading(doc, "3.2 ¿Mejoraron el R² y los residuales?", level=2)
    add_para(doc, "La Tabla 2 compara el modelo de la Tarea 1.6 con el refinado en las dos series.", keep=True)

    def fila(serie, nombre, r):
        pr = r["m"].params
        b0 = f"{pr['const'] * 1000:.2f}" if "const" in pr else "0 (forzado)"
        return [f"{serie}: {nombre}", b0, f"{pr['PM10'] * 1000:.3f}", f"{pr['PM2.5'] * 1000:.3f}",
                f(r["r2_modelo"]), f(r["r2_media"]), f"{r['rmse'] * 1000:.2f}", f"{r['sesgo'] * 1000:+.2f}"]
    make_table(doc, ["Serie y modelo", "b₀ (ppb)", "b PM10", "b PM2.5", "R² del modelo",
                     "R² vs. promedio", "RMSE (ppb)", "Residual medio (ppb)"],
               [fila("Completada", "1.6 con intercepto", c16), fila("Completada", "1.7 por el origen", c17),
                fila("Observada", "1.6 con intercepto", o16), fila("Observada", "1.7 por el origen", o17)],
               "Tabla 2. Modelo de la Tarea 1.6 contra el modelo refinado, datos depurados. Los "
               "coeficientes están en ppb de NO₂ por µg/m³. El R² del modelo es el que reporta "
               "statsmodels; en el modelo por el origen se calcula contra cero.",
               [1.9, 0.7, 0.6, 0.6, 0.7, 0.7, 0.7, 0.9], hl=1)
    add_para(doc,
        f"A primera vista el modelo refinado mejora muchísimo: el R² pasa de {f(c16['r2_modelo'])} a "
        f"{f(c17['r2_modelo'])} en la serie completada y de {f(o16['r2_modelo'])} a "
        f"{f(o17['r2_modelo'])} en la observada. Pero esa mejora no es real. Ese R² está calculado "
        "contra cero, y como el NO₂ casi nunca está cerca de cero, casi cualquier plano que pase "
        "cerca de los puntos da un valor alto. Cuando calculo el R² contra el promedio en los dos "
        f"modelos, el refinado baja a {f(c17['r2_media'])} en la serie completada y a "
        f"{f(o17['r2_media'])} en la observada, es decir, explica menos que el de la 1.6. El RMSE "
        f"también sube, de {c16['rmse'] * 1000:.2f} a {c17['rmse'] * 1000:.2f} ppb en la serie "
        f"completada y de {o16['rmse'] * 1000:.2f} a {o17['rmse'] * 1000:.2f} ppb en la observada.")
    add_para(doc,
        "Con los residuales pasa lo mismo. El modelo de la 1.6 tiene residual medio cero, como debe "
        "ser en una regresión con intercepto. El refinado tiene un residual medio de "
        f"{c17['sesgo'] * 1000:+.2f} ppb en la serie completada: en promedio se queda corto y "
        "subestima el NO₂. La Figura 1 lo muestra comparando lo real con lo estimado, y la Figura 2 "
        "con los residuales.", keep=True)
    add_pic(doc, p_rv, 6.4)
    add_caption(doc, "Figura 1. NO₂ real contra NO₂ estimado por el modelo de la Tarea 1.6 y por el "
                     "modelo refinado, serie completada depurada. La línea punteada es la predicción perfecta.")
    add_pic(doc, p_res, 6.4)
    add_caption(doc, "Figura 2. Residuales de los dos modelos, serie completada depurada. Arriba, "
                     "residuales contra valor estimado con su promedio móvil; abajo, su distribución.")
    add_para(doc,
        "En la parte de arriba de la Figura 2 se ve el problema. En el modelo de la 1.6 el promedio "
        "móvil de los residuales queda pegado al cero en casi todo el rango. En el refinado tiene una "
        "pendiente clara: queda por encima de cero cuando el NO₂ estimado es bajo y por debajo cuando "
        "es alto. O sea, en las horas con pocas partículas el modelo predice casi nada de NO₂ cuando "
        "en realidad sí hay, y en las horas con muchas partículas se pasa. Un residual con tendencia "
        "así indica que al modelo le falta algo, en este caso la constante. (La línea recta de puntos "
        "que se ve arriba en las dos gráficas son las horas imputadas que quedaron fijas en el valor "
        f"límite de {t16.pred15(t16.co_lim):.3f} ppm en la Tarea 1.6.) En el histograma "
        "de abajo la distribución de los residuales del modelo refinado está corrida a la derecha del "
        "cero. Lo que no cambia casi nada es la correlación entre el NO₂ real y el estimado "
        f"({f(c16['r_y_yhat'])} contra {f(c17['r_y_yhat'])}), así que los dos modelos ordenan igual "
        "las horas de más a menos contaminadas; lo que el refinado pierde es el nivel.")

    add_heading(doc, "3.3 ¿Cómo se relacionan las partículas entre sí?", level=2)
    add_para(doc,
        "El PM10 y el PM2.5 son, por mucho, el par más correlacionado de todo el análisis "
        f"(r = {f(rl_o.loc['PM10', 'PM2.5'])} con datos observados y {f(rl_c.loc['PM10', 'PM2.5'])} en la "
        "serie completada, ya sin atípicos). Tiene sentido, porque el PM10 incluye al PM2.5. Una recta "
        f"por el origen lo resume bien: PM2.5 ≈ {O['frac']:.2f}·PM10 (R²₀ = {f(O['frac_r2'])}), o sea "
        f"que alrededor del {O['frac'] * 100:.0f} % de la masa de PM10 son partículas finas. El VIF de "
        f"{f(O['vif'], 2)} en los datos observados confirma que hay colinealidad moderada entre las dos.")
    add_para(doc,
        "Para ver si había distintos tipos de horas según sus partículas, apliqué K-medias a las "
        "horas observadas. La curva del codo (Figura 3, derecha) baja rápido hasta k = 3 y luego se "
        "aplana, así que usé tres grupos. Sus promedios están en la Tabla 3.", keep=True)
    add_pic(doc, p_pm, 6.4)
    add_caption(doc, "Figura 3. Izquierda: PM2.5 contra PM10 con los tres grupos de K-medias (las cruces "
                     "son los centroides) y la recta por el origen. Derecha: J para k de 1 a 7. Horas "
                     "con NO₂ observado, sin atípicos.")
    make_table(doc, ["Grupo", "Horas", "PM10 (µg/m³)", "PM2.5 (µg/m³)", "PM2.5 / PM10", "NO₂ (ppb)"],
               [[NOMBRES_G[g], num(grupos.loc[g, "n"]), f(grupos.loc[g, "PM10"], 1), f(grupos.loc[g, "PM25"], 1),
                 f(grupos.loc[g, "ratio"], 2), f(grupos.loc[g, "NO2"] * 1000, 1)] for g in range(K)],
               "Tabla 3. Promedio de cada grupo de K-medias (k = 3), horas con NO₂ observado.",
               [1.0, 0.8, 1.1, 1.1, 1.0, 0.9])
    add_para(doc,
        "Los grupos no separan tipos de partículas sino niveles de contaminación: bajo, medio y "
        "alto. Los tres caen sobre la misma diagonal, y la proporción de finas cambia poco, de "
        f"{f(grupos.loc[0, 'ratio'], 2)} en el grupo bajo a {f(grupos.loc[2, 'ratio'], 2)} en el alto. Esto refuerza la idea de que el PM10 y el PM2.5 se "
        "mueven casi como una sola variable. Lo interesante es que el NO₂ también sube de grupo en "
        f"grupo, de {f(grupos.loc[0, 'NO2'] * 1000, 1)} a {f(grupos.loc[2, 'NO2'] * 1000, 1)} ppb, aunque "
        "el NO₂ no se usó para formar los grupos. Pero incluso en el grupo bajo el NO₂ promedio es de "
        f"{f(grupos.loc[0, 'NO2'] * 1000, 1)} ppb, lejos de cero, y eso explica por qué forzar el "
        "modelo por el origen no funciona aquí.")

    # ---------------- 4. Interpretacion
    add_heading(doc, "4. ¿Qué nos dice este modelo?")
    p16 = c16["m"].params; p17 = c17["m"].params
    add_para(doc,
        "En el modelo refinado de la serie completada, cada µg/m³ de PM10 se asocia con "
        f"{p17['PM10'] * 1000:.2f} ppb de NO₂ y cada µg/m³ de PM2.5 con {p17['PM2.5'] * 1000:.2f} ppb. "
        "Los dos coeficientes salieron más grandes que en la 1.6 "
        f"({p16['PM10'] * 1000:.2f} y {p16['PM2.5'] * 1000:.2f}). No es que las partículas ahora pesen "
        "más, sino que, al quitarle el intercepto, las pendientes tienen que absorber el NO₂ que antes "
        f"explicaba la constante, que era de {p16['const'] * 1000:.1f} ppb en la serie completada y "
        f"de {o16['m'].params['const'] * 1000:.1f} ppb en la observada, y que era claramente distinta de cero.")
    add_para(doc,
        "Esa constante tiene una lectura física. Con el CO forzar el origen funcionaba, porque los dos "
        "salen del mismo escape de los coches. Con las partículas no pasa lo mismo: hay horas con "
        "poco polvo y poco PM2.5 que siguen teniendo NO₂ por el tráfico, y horas con muchas partículas "
        "por polvo o construcción que no tienen tanto NO₂. Las partículas tienen fuentes que el NO₂ no "
        "comparte, así que no hay razón para que el NO₂ valga cero cuando las partículas valen cero. "
        "El modelo refinado lo deja ver: al quitar la constante, el ajuste empeora y los residuales "
        "quedan sesgados.")

    # ---------------- 5. Conclusiones
    add_heading(doc, "5. Conclusiones")
    add_para(doc,
        "Quitar los atípicos de NO₂, PM10 y PM2.5 bajó la correlación, porque esas horas extremas, en "
        "las que todo sube al mismo tiempo, la estaban inflando. El resultado depurado es más débil, "
        "pero más honesto. Forzar el plano por el origen subió el R² del modelo de "
        f"{f(c16['r2_modelo'])} a {f(c17['r2_modelo'])}, pero es un efecto de cómo se calcula ese R² "
        f"sin intercepto. Medido contra el promedio bajó a {f(c17['r2_media'])}, el RMSE subió y los "
        "residuales quedaron con un sesgo positivo, así que el modelo refinado es peor que el de la 1.6.")
    add_para(doc,
        f"El PM10 y el PM2.5 están muy relacionados (r ≈ {f(rl_o.loc['PM10', 'PM2.5'], 2)}) y el PM2.5 es "
        f"cerca del {O['frac'] * 100:.0f} % del PM10. K-medias no encontró grupos de distinto tipo, solo "
        "niveles de contaminación sobre la misma diagonal, y el NO₂ sube junto con esos niveles. Para "
        "describir el NO₂ con partículas me quedaría con el modelo con intercepto de la Tarea 1.6, o "
        "con la separación en fracción fina y gruesa, y dejaría la recta por el origen para pares "
        "como el NO₂ y el CO, que sí comparten fuente.")

    add_heading(doc, "Referencias")
    add_apa(doc, "Montgomery, D. C., Peck, E. A., & Vining, G. G. (2021). Introduction to linear "
                 "regression analysis (6.ª ed.). Wiley.")
    add_apa(doc, "Ruiz Chávez, C. F. (2026). Visualización de grandes bases de datos: Sesión 9 "
                 "[Diapositivas de clase]. Maestría en Ciencia de los Datos, Centro Universitario de "
                 "Ciencias Económico Administrativas, Universidad de Guadalajara.")
    add_apa(doc, "Secretaría de Medio Ambiente y Desarrollo Territorial. (2024). Sistema de "
                 "Monitoreo Atmosférico de Jalisco (SIMAJ): base de datos de calidad del aire 2024. "
                 "Gobierno del Estado de Jalisco. https://siga.jalisco.gob.mx/aire/")
    add_apa(doc, "Seinfeld, J. H., & Pandis, S. N. (2016). Atmospheric chemistry and physics: From "
                 "air pollution to climate change (3.ª ed.). Wiley.")
    out = OUT + r"\Tarea1_7_VisualizacionGBD_CarlosPulidoRosas.docx"
    doc.save(out); print("DOCX:", out)
    subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", OUT, out],
                   check=True, capture_output=True)
    print("PDF:", out.replace(".docx", ".pdf"))
    return out


if __name__ == "__main__":
    build()
